# -*- coding: utf-8 -*-
"""Extract / apply Azerbaijani CV editable blocks (header, profile, non-publication tables)."""
from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from _normalize_cv_contact_links import set_cell_text

ROOT = Path(__file__).resolve().parents[1]
AZ_CV_DIR = ROOT / "az" / "scientists" / "CV" / "Word"
EN_CV_DIR = ROOT / "en" / "scientists" / "CV" / "Word"

PUBLICATION_HEADINGS = re.compile(
    r"^(nəşrlər|kitablar?\b|kitab fəsilləri|rəylənmiş|peer-reviewed|"
    r"konfrans məqalələri|dəvətli məruzələr|patentlər?\b|proqram təminatı\b)",
    re.I,
)
EDITABLE_TABLE_HEADINGS = re.compile(
    r"^(təhsil|elmi və peşəkar vəzifələr|mükafatlar|"
    r"qrantlar və tədqiqat|elmi göstəricilər|tədqiqat göstəriciləri|"
    r"tədris|elmi rəhbərlik|rəhbərlik|dillər|daab\b)",
    re.I,
)
CONTACT_HEADING = re.compile(r"^əlaqə\b", re.I)
PROFILE_HEADING = re.compile(r"^profil xülasəsi\b", re.I)
SECTION_HEADING = re.compile(
    r"^(əlaqə\b|profil xülasəsi\b|təhsil\b|elmi və peşəkar|"
    r"mükafatlar\b|qrantlar\b|elmi göstəricilər\b|tədqiqat göstəriciləri\b|"
    r"tədris\b|elmi rəhbərlik\b|rəhbərlik\b|dillər\b|daab\b|"
    r"nəşrlər\b|kitablar?\b|kitab fəsilləri\b|rəylənmiş|konfrans|"
    r"dəvətli məruzələr|patentlər?\b|son yenilənmə|date of last update)",
    re.I,
)

META_AZ_RE = re.compile(
    r"\b("
    r"cv-də|cv də|il verilmir|illər verilmir|cv prints|the cv\b|"
    r"qeyd olunmur|qeyd edilmir"
    r")\b",
    re.I,
)


def iter_blocks(doc: Document):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield ("p", Paragraph(child, doc))
        elif child.tag == qn("w:tbl"):
            yield ("t", Table(child, doc))


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def extract_cv(path: Path) -> dict:
    doc = Document(str(path))
    blocks = list(iter_blocks(doc))
    header: list[str] = []
    profile = ""
    tables: list[dict] = []
    in_publications = False
    current_section = ""

    i = 0
    while i < len(blocks):
        kind, obj = blocks[i]
        if kind == "p":
            text = (obj.text or "").strip()
            if CONTACT_HEADING.match(text):
                break
            if text:
                header.append(text)
        i += 1

    while i < len(blocks):
        kind, obj = blocks[i]
        if kind == "p":
            text = (obj.text or "").strip()
            if not text:
                i += 1
                continue
            if SECTION_HEADING.match(text) or PUBLICATION_HEADINGS.match(text):
                current_section = text
                if PUBLICATION_HEADINGS.match(text):
                    in_publications = True
                elif (
                    EDITABLE_TABLE_HEADINGS.match(text)
                    or PROFILE_HEADING.match(text)
                    or CONTACT_HEADING.match(text)
                ):
                    in_publications = False
                if PROFILE_HEADING.match(text):
                    j = i + 1
                    while j < len(blocks):
                        k2, o2 = blocks[j]
                        if k2 == "p":
                            t2 = (o2.text or "").strip()
                            if t2 and not SECTION_HEADING.match(t2):
                                profile = t2
                                break
                            if t2 and SECTION_HEADING.match(t2):
                                break
                        elif k2 == "t":
                            break
                        j += 1
            i += 1
            continue

        if in_publications or CONTACT_HEADING.match(current_section):
            i += 1
            continue
        if not EDITABLE_TABLE_HEADINGS.match(current_section):
            i += 1
            continue

        rows = [[c.text.strip() for c in row.cells] for row in obj.rows]
        tables.append({"section": current_section, "rows": rows})
        i += 1

    return {
        "file": path.name,
        "header": header,
        "profile": profile,
        "tables": tables,
        "has_meta": bool(META_AZ_RE.search(profile) or any(
            META_AZ_RE.search(" ".join(r)) for t in tables for r in t["rows"]
        )),
    }


def apply_cv(path: Path, edit: dict) -> list[str]:
    doc = Document(str(path))
    blocks = list(iter_blocks(doc))
    changes: list[str] = []
    new_header = edit.get("header")
    new_profile = edit.get("profile")
    new_tables = edit.get("tables")

    hi = 0
    i = 0
    while i < len(blocks):
        kind, obj = blocks[i]
        if kind == "p":
            text = (obj.text or "").strip()
            if CONTACT_HEADING.match(text):
                break
            if text and new_header is not None and hi < len(new_header):
                if text != new_header[hi]:
                    set_paragraph_text(obj, new_header[hi])
                    changes.append(f"header[{hi}]")
                hi += 1
        i += 1

    in_publications = False
    current_section = ""
    table_idx = 0

    while i < len(blocks):
        kind, obj = blocks[i]
        if kind == "p":
            text = (obj.text or "").strip()
            if text and (SECTION_HEADING.match(text) or PUBLICATION_HEADINGS.match(text)):
                current_section = text
                if PUBLICATION_HEADINGS.match(text):
                    in_publications = True
                elif (
                    EDITABLE_TABLE_HEADINGS.match(text)
                    or PROFILE_HEADING.match(text)
                    or CONTACT_HEADING.match(text)
                ):
                    in_publications = False
                if PROFILE_HEADING.match(text) and new_profile is not None:
                    j = i + 1
                    while j < len(blocks):
                        k2, o2 = blocks[j]
                        if k2 == "p":
                            t2 = (o2.text or "").strip()
                            if t2 and not SECTION_HEADING.match(t2):
                                if t2 != new_profile:
                                    set_paragraph_text(o2, new_profile)
                                    changes.append("profile")
                                break
                            if t2 and SECTION_HEADING.match(t2):
                                break
                        elif k2 == "t":
                            break
                        j += 1
            i += 1
            continue

        if in_publications or CONTACT_HEADING.match(current_section):
            i += 1
            continue
        if not EDITABLE_TABLE_HEADINGS.match(current_section):
            i += 1
            continue

        if new_tables is not None and table_idx < len(new_tables):
            src_rows = new_tables[table_idx].get("rows") or []
            # Resize rows if needed
            while len(obj.rows) < len(src_rows):
                _clone_row(obj, 0 if obj.rows else -1)
            for r_i, row in enumerate(obj.rows):
                if r_i >= len(src_rows):
                    break
                cells_src = src_rows[r_i]
                for c_i, cell in enumerate(row.cells):
                    if c_i >= len(cells_src):
                        break
                    old = cell.text.strip()
                    new = cells_src[c_i]
                    if old != new:
                        set_cell_text(cell, new)
                        changes.append(f"table[{table_idx}].r{r_i}c{c_i}")
            table_idx += 1
        i += 1

    if changes:
        _save(path, doc)
    return changes


def _clone_row(table: Table, row_idx: int = 0):
    if not table.rows:
        return None
    row = table.rows[max(0, row_idx)]
    new_tr = deepcopy(row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _save(path: Path, doc: Document) -> None:
    tmp = path.with_name(f"__tmp_{path.name}")
    try:
        doc.save(str(tmp))
        tmp.replace(path)
    except OSError:
        alt = path.with_name(f"__tmp2_{path.name}")
        doc.save(str(alt))
        alt.replace(path)


def ensure_awards_table(doc: Document) -> Table:
    blocks = list(iter_blocks(doc))
    for i, (kind, obj) in enumerate(blocks):
        if kind == "p" and (obj.text or "").strip().lower().startswith("mükafatlar"):
            if i + 1 < len(blocks) and blocks[i + 1][0] == "t":
                return blocks[i + 1][1]

    insert_before = None
    for kind, obj in blocks:
        if kind != "p":
            continue
        text = (obj.text or "").strip()
        if PUBLICATION_HEADINGS.match(text) or text.lower().startswith("son yenilənmə"):
            insert_before = obj._p
            break

    body = doc.element.body
    h = doc.add_paragraph("Mükafatlar, fəxri adlar və təqaüdlər")
    h_el = h._p
    body.remove(h_el)
    if insert_before is not None:
        insert_before.addprevious(h_el)
    else:
        body.append(h_el)

    table = doc.add_table(rows=1, cols=3)
    tbl_el = table._tbl
    body.remove(tbl_el)
    h_el.addnext(tbl_el)
    set_cell_text(table.rows[0].cells[0], "İl")
    set_cell_text(table.rows[0].cells[1], "Mükafat / fəxri ad")
    set_cell_text(table.rows[0].cells[2], "Təltif edən qurum")
    return table


def extract_all(out_path: Path | None = None) -> list[dict]:
    items = [extract_cv(p) for p in sorted(AZ_CV_DIR.glob("*_CV_AZ.docx"))]
    if out_path:
        out_path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    return items
