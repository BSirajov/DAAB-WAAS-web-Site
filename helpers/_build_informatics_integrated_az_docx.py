"""Build the integrated Azerbaijani Informatics topic guide (DOCX).

Sources (read-only):
- documents/Popular_Topics_in_Informatics_Reworked_AZ.docx
- az/complex-topics-informatics.html

Output:
- documents/Popular_Topics_in_Informatics_Integrated_AZ.docx
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from daab_docx_export import (  # noqa: E402
    BLUE,
    FONT_BODY,
    FONT_HEADING,
    GOLD,
    INK,
    MUTED,
    NAVY,
    QUOTE_FILL,
    add_code_block,
    add_formatted_run,
    add_hyperlink,
    add_page_number_field,
    add_table,
    set_update_fields_on_open,
    setup_styles,
    shade_paragraph,
)

OUT = ROOT / "documents" / "Popular_Topics_in_Informatics_Integrated_AZ.docx"
NOTE_FILL = "FFF8E8"
EXAMPLE_FILL = "F3F9FD"
SITE_AZ = "https://daab-waas.com/az"


def _set_run_font(run, size=9, color=MUTED):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = FONT_BODY


def setup_az_header_footer(document: Document, header_title: str) -> None:
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    header = section.header
    header._element.clear()
    hp = header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run(header_title)
    _set_run_font(r)

    section.first_page_header._element.clear()
    section.first_page_header.add_paragraph()

    footer = section.footer
    footer._element.clear()
    fp = footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("DAAB / WAAS   |   Səhifə ")
    _set_run_font(r1)
    add_page_number_field(fp, align_center=False)
    r2 = fp.add_run("   |   sentyabr 2026")
    _set_run_font(r2)

    first_footer = section.first_page_footer
    first_footer._element.clear()
    ff = first_footer.add_paragraph()
    ff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = ff.add_run("daab-waas.com")
    _set_run_font(rf, color=GOLD)


def add_title_page(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()

    org = document.add_paragraph()
    org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = org.add_run("Dünya Azərbaycanlı Alimlər Birliyi")
    r.font.name = FONT_HEADING
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED

    org2 = document.add_paragraph()
    org2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = org2.add_run("World Association of Azerbaijani Scientists")
    r2.font.size = Pt(11)
    r2.font.italic = True
    r2.font.color.rgb = MUTED

    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = title.add_run("İnformatika üzrə aktual mövzular\nvə onların aydın izahı")
    rt.font.name = FONT_HEADING
    rt.font.size = Pt(26)
    rt.font.bold = True
    rt.font.color.rgb = NAVY

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = rule.add_run("━" * 36)
    rr.font.color.rgb = GOLD

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run(
        "İnteqrasiya olunmuş bələdçi\n"
        "«Çətin mövzu, aydın izah» müsabiqəsi üçün mövzu seçimi"
    )
    rs.font.size = Pt(13)
    rs.font.italic = True
    rs.font.color.rgb = BLUE

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(f"Hazırlanma tarixi: {datetime.now().strftime('%d.%m.%Y')}")
    rm.font.size = Pt(10)
    rm.font.color.rgb = INK

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run(
        "Bu sənəd iki mənbənin birləşdirilmiş nəşridir.\n"
        "Əsas mənbələr dəyişdirilmədən saxlanılıb."
    )
    rn.font.size = Pt(10)
    rn.font.color.rgb = MUTED

    document.add_page_break()


def add_toc(document: Document) -> None:
    document.add_heading("Mündəricat", level=1)
    note = document.add_paragraph()
    add_formatted_run(
        note,
        "Microsoft Word-də aşağıdakı sahənin üzərinə klikləyin və **F9** düyməsini basın, "
        "və ya sağ klik → **Sahəni yenilə** → **Bütün cədvəli yenilə**. "
        "Keçidlər başlıqların 1–3-cü səviyyələrini əhatə edir.",
    )
    toc_p = document.add_paragraph()
    add_toc_field_az(toc_p)
    document.add_page_break()


def p(document: Document, text: str, *, space_after: int = 8) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    add_formatted_run(para, text)


def note_box(document: Document, text: str, fill: str = NOTE_FILL) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    shade_paragraph(para, fill)
    add_formatted_run(para, text)


def example_box(document: Document, text: str) -> None:
    note_box(document, text, EXAMPLE_FILL)


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        para = document.add_paragraph(style="List Bullet")
        add_formatted_run(para, item)


def numbered(document: Document, items: list[str]) -> None:
    for item in items:
        para = document.add_paragraph(style="List Number")
        add_formatted_run(para, item)


def add_toc_field_az(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    run2._r.append(instr)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)

    run4 = paragraph.add_run(
        "Word-də açın, sonra F9 düyməsini basın və ya sahəni yeniləyin."
    )
    run4.font.italic = True
    run4.font.color.rgb = MUTED
    run4.font.size = Pt(10)

    run5 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run5._r.append(fld_end)


def write_intro(document: Document) -> None:
    document.add_heading("1. Bələdçidən istifadə", level=1)
    p(
        document,
        "Dünya Azərbaycanlı Alimlər Birliyinin (DAAB) bu bələdçisi informatika üzrə "
        "tədris materialı hazırlamaq istəyən müəllimlərə, mütəxəssislərə və "
        "«Çətin mövzu, aydın izah» müsabiqəsinin iştirakçılarına mövzu seçməkdə kömək edir. "
        "Burada təhsil baxımından əhəmiyyətli və praktikada aktual olan mövzular, "
        "onların öyrənilməsindəki çətinliklər və aydın izah üçün nümunələr bir araya gətirilir.",
    )
    p(
        document,
        "Sənəd işçi bələdçidir. Siyahı mümkün mövzuları məhdudlaşdırmır. "
        "İştirakçılar göstərilən istiqamətlərdən birini seçə və ya təhsil əhəmiyyətini "
        "əsaslandıraraq yeni mövzu təklif edə bilərlər. Materialın dəyəri auditoriyanın "
        "konkret sualına cavab verməsi və öyrənənə yeni bilik və ya bacarıq qazandırması ilə "
        "müəyyən olunur.",
    )
    p(
        document,
        "Mövzu seçərkən üç məsələni birlikdə nəzərə almaq faydalıdır: mövzunun öyrənən üçün "
        "əhəmiyyəti, tətbiq sahəsi və başa düşülməsində yaranan çətinlik. "
        "İnformatika mövzularının populyarlığı və çətinliyi üzrə hamı tərəfindən qəbul edilmiş "
        "vahid qlobal reytinq yoxdur. Bu sənəddə «maraq» və ya «populyarlıq» öyrənənlərin "
        "ehtiyacları, tədris proqramlarının prioritetləri və texnologiya sahəsindəki fəaliyyət "
        "barədə göstəricilərin ümumi şərhidir; internet axtarışlarının sayını bildirmir.",
    )
    p(
        document,
        "Otuz mövzudan ibarət cədvəldəki sıra və «yüksək», «çox yüksək» kimi qiymətlər "
        "istiqamətləndirici qiymətləndirmələrdir. Bunlar vahid ölçmə üsulu ilə hesablanmış "
        "ballar deyil. Mövzunun çətinliyi auditoriyanın yaşından, ilkin biliyindən, seçilmiş "
        "izah səviyyəsindən və istifadə olunan nümunələrdən asılıdır.",
    )

    document.add_heading("1.1 Bu inteqrasiya olunmuş nəşr necə qurulub", level=2)
    p(
        document,
        "Bələdçi iki mövcud mənbənin məzmununu birləşdirir. Hər iki orijinal mənbə "
        "layihədə olduğu kimi saxlanılıb; bu fayl onların yerini tutmur.",
    )
    bullets(
        document,
        [
            "**Yenidən işlənmiş Word sənədi** — `documents/Popular_Topics_in_Informatics_Reworked_AZ.docx`. "
            "Mənbələrin ehtiyatlı şərhi, termin izahları, mövzu təklifi üçün beş sual və "
            "nömrələnmiş istinadlar buradadır.",
            "**Azərbaycanca veb səhifə** — `az/complex-topics-informatics.html`. "
            "Məktəb üçün 23, universitet üçün 25 nömrələnmiş mövzu, hər bənd üzrə qısa izah "
            "və nümunə, həmçinin müsabiqənin əlaqəli səhifələrinə keçidlər buradadır.",
        ],
    )
    p(
        document,
        "Mənbələr bir-birini tamamlayanda məlumat eyni bölmədə birləşdirilib. "
        "Eyni mövzunun iki nümunəsi varsa, hər ikisi saxlanılıb. "
        "Mənbələr fərqli söz, qruplaşma və ya qətiyyət seçəndə seçim gizlədilməyib: "
        "fərq 2-ci bölmədə göstərilib. Məktəb və universitet bölmələrində əvvəlcə "
        "nömrələnmiş mövzu siyahısı, sonra eyni sırada izah və nümunələr gəlir.",
    )
    p(
        document,
        "Əvvəlcə 3-cü bölmədə mənbələrin nəyi göstərdiyini oxuyun. Sonra 4-cü bölmədən "
        "ümumi prioritetə baxın, 5-ci və 6-cı bölmələrdən auditoriyaya uyğun altmövzuları "
        "seçin. 7-ci bölmə çətinlikləri izah edir, 8-ci bölmə ilk nəşrlər üçün plan verir, "
        "9-cu bölmə isə mövzu təklifini hazırlamağa kömək edir. Terminlərin qısa izahı və "
        "birləşdirilmiş mənbələr sənədin sonundadır.",
    )


def write_discrepancies(document: Document) -> None:
    document.add_heading("2. Mənbələr arasındakı fərqlər", level=1)
    p(
        document,
        "Aşağıdakı cədvəl iki mənbənin eyni fakta fərqli söz, qruplaşma və ya qətiyyətlə "
        "yanaşdığı yerləri toplayır. Bu, səhv və ya düzgün versiya seçimi deyil. "
        "İnteqrasiya olunmuş mətndə hər iki ifadə saxlanılıb; işlək izahda isə daha ehtiyatlı "
        "və ya daha dəqiq olan ifadə əsas götürülüb, o biri qısa qeyd kimi göstərilib.",
    )
    add_table(
        document,
        [
            ["Məsələ", "Yenidən işlənmiş sənəd", "Veb səhifə", "Bu bələdçidə necə göstərilib"],
            [
                "Python tövsiyəsi",
                "Praktik başlanğıc seçimi ola bilər; vahid «ən yaxşı dil» hökmü deyil.",
                "Hazırda proqramlaşdırma, verilənlər elmi və süni intellekt üçün ən uyğun ümumi tədris dilidir.",
                "Eyni istiqamət, fərqli qətiyyət. Hər iki ifadə 8.2-ci bölmədə saxlanılıb.",
            ],
            [
                "TypeScript göstəricisi",
                "2025-ci ilin avqustunda GitHub-un istifadə göstəricisinə görə Python və JavaScript-i keçib.",
                "2025-ci ildə ən çox istifadə olunan dil olub; platforma dəqiqləşdirilmir.",
                "GitHub Octoverse 2025 göstəricisi əsas götürülür; veb səhifənin daha qısa ifadəsi qeyd olunur.",
            ],
            [
                "CSTA «ixtisaslar»",
                "Ali təhsil ixtisasları deyil, məktəb səviyyəsində ixtisaslaşma istiqamətləridir.",
                "Süni intellekt, kibertəhlükəsizlik və s. üzrə ixtisaslar müəyyən edir.",
                "Veb səhifənin siyahısı saxlanılıb; yenidən işlənmiş sənədin ehtiyatı əlavə edilib.",
            ],
            [
                "CSTA-da süni intellekt",
                "Baza səviyyəsində istiqamətlərə inteqrasiya olunur; yuxarı siniflərdə dərinləşir.",
                "Standartlar süni intellekti prioritet sayır.",
                "Hər iki vurğu 3.2-ci bölmədə yan-yana verilib.",
            ],
            [
                "Universitet mövzularının qruplaşması",
                "Yeddi tematik qrup. Paradiqmalar və kompilyatorlar icra ilə, avtomatlar isə alqoritmlərlə birlikdədir.",
                "25 nömrələnmiş bənd. Nəzəriyyə ayrıca 25-ci bənddir; paradiqmalar 23–24-cü bəndlərdir.",
                "Nömrələnmiş siyahı veb səhifədəki 25 bənddir. Qrup fərqi 6-cı bölmənin əvvəlində qeyd olunur.",
            ],
            [
                "Prioritet cədvəlinin başlıqları",
                "Qısa mövzu adları, məsələn «böyük dil modelləri».",
                "Bəzən sual forması, məsələn «iri dil modelləri necə işləyir».",
                "Cədvəldə yenidən işlənmiş adlar verilib; veb səhifənin variantı eyni sətirdə göstərilib.",
            ],
            [
                "Auditoriya adı",
                "Məktəblilər və universitet tələbələri.",
                "Şagirdlər və orta məktəb şagirdləri; kollec və universitet tələbələri.",
                "Hər iki auditoriya adı saxlanılıb. «Kollec» yalnız veb səhifədə keçir.",
            ],
            [
                "Deadlock və şəbəkə tıxacı",
                "Qarşılıqlı kilidlənmə; şəbəkədə sıxlıq.",
                "Tıxac (deadlock); tıxaclıq (congestion).",
                "Eyni kök iki fərqli anlayışı qarışdıra bilər. Əsas izahda hər iki cüt saxlanılıb.",
            ],
            [
                "Heap",
                "«Dinamik yaddaş»; heap həm yaddaş bölgəsi, həm də ayrıca struktur ola bilər.",
                "Cədvəldə «heap»; izahda fərq açılmır.",
                "İkimənalılıq 6-cı və 11-ci bölmələrdə izah olunur.",
            ],
            [
                "Termin cütləri",
                "Polimorfizm; iki amilli autentifikasiya; Big-O; uyğunluq; izolyasiya; orkestrasiya; konsensus; obyektyönlü.",
                "Çoxformalıq; iki faktorlu təsdiq; Böyük-O; tutarlılıq; izolə; orkestrləşdirmə; razılaşma; obyektyönümlü.",
                "Sinonimlər və orfoqrafiya variantlarıdır. İlk dəfə keçəndə hər iki forma verilir.",
            ],
            [
                "Müsabiqə keçidləri",
                "Yerli işçi ünvan: http://127.0.0.1:8010/az/complex-topics.html",
                "Saytdakı dəvət, təşkilat və təqdimat səhifələri.",
                "İctimai keçidlər əsas götürülür; yerli ünvan inkişaf nüsxəsi kimi qeyd olunur.",
            ],
        ],
    )
    p(
        document,
        "Bundan başqa, veb səhifədə universitet mövzularının 11, 23 və 24-cü bəndlərində "
        "ayrıca nümunə yoxdur. Yenidən işlənmiş sənəddəki qrup izahları və əlavə nümunələr "
        "həmin boşluğu doldurur. Məktəb mövzularında isə veb səhifənin bənd-bənd nümunələri "
        "yenidən işlənmiş sənədin qrup izahlarını tamamlayır.",
    )


def write_sources(document: Document) -> None:
    document.add_heading("3. Mövzu seçiminin əsaslandığı mənbələr", level=1)
    document.add_heading("3.1 Mənbələrin birgə istifadəsi", level=2)
    p(
        document,
        "Mövzuların seçimi üç mənbə qrupuna əsaslanır. Hər qrup fərqli suala cavab verir "
        "və digərlərini tamamlayır. Təhsil standartı bir mövzunun öyrədilməsinin vacibliyini "
        "göstərə bilər, lakin onun nə qədər populyar və ya çətin olduğunu birbaşa ölçmür. "
        "Sorğular və platforma hesabatları praktik marağı göstərir. Tədris tədqiqatları isə "
        "öyrənənlərin hansı anlayışlarda çətinlik çəkdiyini müəyyənləşdirməyə kömək edir.",
    )
    add_table(
        document,
        [
            ["Mənbə qrupu", "Əsas sual", "Mövzu seçiminə töhfəsi"],
            [
                "Təhsil standartları və proqram tövsiyələri",
                "Nəyi bilmək və bacarmaq lazımdır?",
                "Əsas anlayışları və uyğun təhsil səviyyəsini göstərir.",
            ],
            [
                "Peşəkar fəaliyyət və maraq göstəriciləri",
                "Hansı texnologiyalar istifadə olunur və diqqət çəkir?",
                "Mövzunun müasir tətbiqlərlə əlaqəsini göstərir.",
            ],
            [
                "Öyrənmə və tədris tədqiqatları",
                "Öyrənənlər harada və niyə çətinlik çəkir?",
                "Aydın izaha ehtiyacı olan anlayışları müəyyənləşdirməyə kömək edir.",
            ],
        ],
    )

    document.add_heading("3.2 Məktəb təhsili üçün CSTA standartları", level=2)
    p(
        document,
        "**CSTA** — Computer Science Teachers Association, yəni İnformatika Müəllimləri "
        "Assosiasiyasıdır. **PK–12** məktəbəqədər mərhələdən 12-ci sinfə qədər olan təhsil "
        "pillələrini ifadə edir. Burada PK «pre-kindergarten», K isə «kindergarten» "
        "sözlərinin qısaltmasıdır. Bu bölgü Azərbaycan təhsil sisteminin sinifləri ilə "
        "avtomatik eyniləşdirilməməlidir. [1]",
    )
    p(
        document,
        "CSTA-nın 2026-cı il standartları məktəb informatikasını beş əsas istiqamət ətrafında qurur:",
    )
    bullets(
        document,
        [
            "Alqoritmlər və layihələndirmə;",
            "Proqramlaşdırma;",
            "Verilənlər və təhlil;",
            "Sistemlər və təhlükəsizlik;",
            "Hesablama və cəmiyyət.",
        ],
    )
    p(
        document,
        "Süni intellekt baza səviyyəsində bu istiqamətlərə inteqrasiya olunur. "
        "Veb səhifə eyni standartların süni intellekti prioritet saydığını da vurğulayır. "
        "Bu iki ifadə ziddiyyət deyil: inteqrasiya baza səviyyəsini, prioritet isə "
        "standartların verdiyi vurğunu təsvir edir. [1]",
    )
    p(
        document,
        "Yuxarı siniflərdə dərinləşdirilmiş öyrənmə üçün süni intellekt, kibertəhlükəsizlik, "
        "verilənlər elmi, oyunların hazırlanması, fiziki hesablama və proqram təminatının "
        "hazırlanması nəzərdə tutulur. Veb səhifə bunları «ixtisaslar» adlandırır. "
        "Yenidən işlənmiş sənəd isə xəbərdarlıq edir: bunlar ali təhsil ixtisasları deyil, "
        "məktəb səviyyəsində ixtisaslaşma istiqamətləridir.",
    )
    p(
        document,
        "**Fiziki hesablama** proqramın sensorlar və elektron qurğular vasitəsilə real mühitlə "
        "qarşılıqlı işləməsidir. Məsələn, torpağın rütubətini ölçüb suvarmanı işə salan qurğu "
        "bu istiqamətə uyğun layihədir.",
    )

    document.add_heading("3.3 Universitet təhsili üçün CS2023 tövsiyələri", level=2)
    p(
        document,
        "**CS2023** — Computer Science Curricula 2023, yəni kompüter elmləri üzrə tədris "
        "proqramı tövsiyələridir. Sənəd üç peşəkar təşkilatın birgə işidir. Təşkilat adları "
        "aşağıdakı cədvəldə açılır.",
    )
    add_table(
        document,
        [
            ["İxtisar", "Tam adı və izahı"],
            [
                "ACM",
                "Association for Computing Machinery — kompüter elmləri və texnologiyaları üzrə elmi-peşəkar assosiasiya.",
            ],
            [
                "IEEE-CS",
                "IEEE Computer Society — Elektrik və Elektronika Mühəndisləri İnstitutunun kompüter sahəsi üzrə cəmiyyəti. IEEE: Institute of Electrical and Electronics Engineers.",
            ],
            [
                "AAAI",
                "Association for the Advancement of Artificial Intelligence — Süni İntellektin İnkişafı Assosiasiyası.",
            ],
        ],
    )
    p(
        document,
        "CS2023 əsasən bakalavr təhsilinə yönəlib. Adında «2023» olsa da, yekun sənəd "
        "tərəfdaş təşkilatlar tərəfindən 2024-cü ildə təsdiqlənib. O, alqoritmlər, süni "
        "intellekt, verilənlərin idarə edilməsi, şəbəkələr, əməliyyat sistemləri, təhlükəsizlik "
        "və proqram mühəndisliyi daxil olmaqla 17 bilik sahəsini əhatə edir. [2, 23]",
    )
    p(
        document,
        "Bu yanaşmada biliklərin tətbiqi, əməkdaşlıq, məsuliyyət və peşəkar davranış da "
        "nəzərə alınır. Müsabiqə materialını planlaşdırarkən «Hansı məlumatı təqdim edirəm?» "
        "sualı ilə yanaşı «Öyrənən bu materialdan sonra nə edə biləcək?» sualını vermək "
        "buna görə faydalıdır. [24]",
    )

    document.add_heading("3.4 AP Computer Science Principles kursu", level=2)
    p(
        document,
        "**AP** — Advanced Placement, məktəblilər üçün universitetin giriş səviyyəsinə uyğun "
        "kurs və imtahan proqramıdır. **CSP** — Computer Science Principles, yəni "
        "«Kompüter elmlərinin əsasları» deməkdir. Proqramı College Board təşkilatı təqdim edir. [3]",
    )
    p(document, "AP CSP beş əsas mövzunu birləşdirir:")
    bullets(
        document,
        [
            "Yaradıcı həllərin hazırlanması;",
            "Verilənlər;",
            "Alqoritmlər və proqramlaşdırma;",
            "Kompüter sistemləri və şəbəkələr;",
            "Hesablama texnologiyalarının təsiri.",
        ],
    )
    p(
        document,
        "Kurs proqramlaşdırma ilə yanaşı internetin işləməsi, verilənlərdən nəticə çıxarılması "
        "və texnologiyanın cəmiyyətə təsiri barədə ilkin anlayış yaradır. CSTA məktəb "
        "təhsilini, AP CSP universitetə giriş səviyyəli hazırlığı, CS2023 isə bakalavr "
        "təhsilini istiqamətləndirir. Bu fərq mövzunun dərinliyini seçmək üçün əhəmiyyətlidir. [1–3]",
    )

    document.add_heading("3.5 Stack Overflow və GitHub hesabatları", level=2)
    p(
        document,
        "**Stack Overflow** proqramlaşdırma suallarının müzakirə olunduğu platformadır. "
        "Onun Developer Survey hesabatı proqramçıların alətlərdən istifadəsi, iş üsulları və "
        "münasibətləri barədə illik sorğudur. 2025-ci ilin nəticələrində 177 ölkədən "
        "49 009 cavab istifadə olunub. İştirakçılar əsasən platformanın öz kanalları "
        "vasitəsilə cəlb edildiyindən nəticələr bütün proqramçıların tam təmsilçi mənzərəsi "
        "sayılmamalıdır. [4, 25]",
    )
    p(
        document,
        "Süni intellektdən istifadə barədə suala cavab verənlərin 84 faizi bu alətlərdən "
        "istifadə etdiyini və ya istifadə etməyi planlaşdırdığını bildirib. Dəqiqliyə etibar "
        "barədə ayrıca sualda 46 faiz etimadsızlıq, 33 faiz isə etimad ifadə edib. "
        "84 faiz göstəricisi yalnız hazırda istifadə edənləri bildirmir. [26]",
    )
    p(
        document,
        "Bu nəticələrdən «Süni intellektin yaratdığı kodu necə yoxlamaq olar?» və "
        "«İnandırıcı cavab niyə səhv ola bilər?» kimi mövzu təklifləri çıxarmaq mümkündür. "
        "Bunlar sorğunun birbaşa tövsiyələri deyil, nəticələrin tədris məqsədilə şərhidir.",
    )
    p(
        document,
        "**GitHub** proqram kodunun saxlanması və birgə hazırlanması platformadır. "
        "**Octoverse** onun illik fəaliyyət hesabatının adıdır, ixtisar deyil. "
        "2025-ci il hesabatı süni intellekt layihələrinin artmasını və həmin ilin avqustunda "
        "TypeScript-in GitHub-un istifadə göstəricisinə görə Python və JavaScript-i keçməsini "
        "vurğulayır. Hesabatda Python-un yeni süni intellekt repozitoriyalarının təxminən "
        "yarısında istifadə olunduğu da göstərilir. **Repozitoriya** layihənin kodunu və "
        "dəyişiklik tarixçəsini saxlayan mühitdir. [5]",
    )
    note_box(
        document,
        "**Fərq.** Veb səhifə TypeScript-in 2025-ci ildə «ən çox istifadə olunan dil» "
        "olduğunu yazır, lakin bunun GitHub göstəricisi olduğunu dəqiqləşdirmir. "
        "Yenidən işlənmiş sənəd platformanı və ayı göstərir. Hər iki mənbə eyni hesabata "
        "istinad edir; bu bələdçi daha dəqiq olan GitHub Octoverse 2025 ifadəsindən istifadə edir.",
    )
    p(
        document,
        "Bu göstəricilər platformadakı fəaliyyəti təsvir edir. Onlar bütün əmək bazarının "
        "tələbatını, mövzunun təhsil dəyərini və ya öyrənmə çətinliyini birbaşa ölçmür.",
    )

    document.add_heading("3.6 Öyrənmə üzrə empirik tədqiqatlar", level=2)
    p(
        document,
        "**Empirik tədqiqat** real müşahidələrə və toplanmış məlumatlara əsaslanan "
        "araşdırmadır. Tədrisdə bu, öyrənənlərlə müsahibələr, tapşırıq cavablarının təhlili "
        "və dərsdən əvvəl və sonra aparılan ölçmələr ola bilər.",
    )
    p(
        document,
        "**Konsepsiya inventarı** (*concept inventory*) əsas anlayışların nə dərəcədə başa "
        "düşüldüyünü və yanlış təsəvvürləri müəyyən edən diaqnostik testdir. Bu testlərdə "
        "səhv cavab variantları da düşüncə səhvlərini üzə çıxarmaq üçün hazırlanır. "
        "Dinamik proqramlaşdırma üzrə belə bir alətin hazırlanmasında əvvəlki tədqiqatlarda "
        "müəyyən olunmuş yanlış təsəvvürlərdən istifadə edilib. [27]",
    )
    p(
        document,
        "Nümunə olaraq aşağıdakı əmrləri ardıcıllıqla nəzərdən keçirək:",
    )
    add_code_block(document, ["a = 5", "b = a", "a = 8"])
    p(
        document,
        "Sonda `b`-nin qiyməti 5-dir. `b`-yə `a`-nın həmin andakı qiyməti mənimsədilib; "
        "sonrakı əmr yalnız `a`-nı dəyişir. «`b` də 8 olar» cavabı mənimsətmənin daim "
        "yenilənən əlaqə kimi başa düşüldüyünü göstərə bilər.",
    )
    p(
        document,
        "Testin yaşa, dilə və öyrədilən mövzulara uyğunluğu yoxlanmalıdır. **SCS1** — "
        "Second CS1, giriş proqramlaşdırmasında anlayışları ölçən qiymətləndirmə vasitəsidir; "
        "**CS1** ilkin kompüter elmləri və ya proqramlaşdırma kursunu bildirir. SCS1 üzrə "
        "tədqiqat alətin tətbiq sərhədlərini və öyrənmə artımını qiymətləndirmək üçün əvvəl "
        "və sonra ölçmələrinin əhəmiyyətini vurğulayır. [28]",
    )


def write_priority(document: Document) -> None:
    document.add_heading("4. Tövsiyə olunan 30 əsas mövzu", level=1)
    p(
        document,
        "Bu cədvəl mövzulara ümumi baxış verir. Sıra ilkin prioritet təklifidir, statistik "
        "reytinq deyil. «Hər iki qrup» məktəbliləri və universitet tələbələrini bildirir; "
        "eyni mövzu bu qruplara fərqli dərinlikdə izah olunmalıdır. «Yuxarı siniflər» və "
        "«universitet» bölgüsü də sərt qəbul şərti kimi nəzərdə tutulmur. Veb səhifə bəzi "
        "sətirlərdə «kollec» də əlavə edir.",
    )
    p(
        document,
        "Cədvəlin «Mövzu» sütunu yenidən işlənmiş sənədin qısa adını verir. "
        "Veb səhifə eyni mövzunu bəzən sual və ya «necə işləyir» formasında yazır; "
        "həmin variant mötərizədə göstərilib, yalnız fərqli olanda.",
    )
    add_table(
        document,
        [
            ["Sıra", "Mövzu", "Əsas auditoriya", "Maraq", "İzah çətinliyi"],
            [
                "1",
                "Generativ süni intellekt və böyük dil modelləri (veb: iri dil modelləri necə işləyir)",
                "Hər iki qrup",
                "Çox yüksək",
                "Çox yüksək",
            ],
            ["2", "Alqoritmlər və hesablama vasitəsilə məsələ həlli", "Hər iki qrup", "Çox yüksək", "Çox yüksək"],
            [
                "3",
                "Dəyişənlər, mənimsətmə, verilənlərin tipləri və proqramın vəziyyəti",
                "Hər iki qrup",
                "Yüksək",
                "Çox yüksək",
            ],
            [
                "4",
                "Səhvlərin tapılması, testlər və xəta mesajlarının oxunması (veb: sazlama, yoxlama)",
                "Hər iki qrup",
                "Çox yüksək",
                "Çox yüksək",
            ],
            ["5", "Kibertəhlükəsizlik, məxfilik və təhlükəsiz rəqəmsal davranış", "Hər iki qrup", "Çox yüksək", "Çox yüksək"],
            [
                "6",
                "Verilənlər elmi, verilənlərin şərhi və vizuallaşdırılması",
                "Hər iki qrup",
                "Çox yüksək",
                "Yüksək",
            ],
            ["7", "Şərtlər, Bul məntiqi və qərarvermə", "Hər iki qrup", "Yüksək", "Yüksək"],
            ["8", "Dövrlər, təkrarlanma və dayanma (veb: bitmə)", "Hər iki qrup", "Yüksək", "Çox yüksək"],
            [
                "9",
                "Funksiyalar, parametrlər, qaytarılan qiymətlər və görünmə sahəsi (veb: əhatə dairəsi)",
                "Hər iki qrup",
                "Yüksək",
                "Çox yüksək",
            ],
            ["10", "İnternetin və Vebin işləməsi", "Hər iki qrup", "Çox yüksək", "Yüksək"],
            [
                "11",
                "Maşın öyrənməsi, təlim verilənləri, modellər və proqnozlar",
                "Yuxarı siniflər və universitet",
                "Çox yüksək",
                "Çox yüksək",
            ],
            ["12", "Verilənlər strukturları və tətbiqləri", "Universitet", "Yüksək", "Çox yüksək"],
            [
                "13",
                "Alqoritmlərin səmərəliliyi və Big-O (veb: Böyük-O yazılışı)",
                "Universitet",
                "Yüksək",
                "Çox yüksək",
            ],
            [
                "14",
                "Verilənlər bazaları, modelləşdirmə və SQL",
                "Yuxarı siniflər və universitet",
                "Yüksək",
                "Çox yüksək",
            ],
            [
                "15",
                "Obyektyönlü proqramlaşdırma (veb: obyektyönümlü)",
                "Yuxarı siniflər və universitet",
                "Yüksək",
                "Çox yüksək",
            ],
            ["16", "Rekursiya", "Yuxarı siniflər və universitet", "Orta–yüksək", "Çox yüksək"],
            [
                "17",
                "Yaddaş, istinadlar, göstəricilər, stek və dinamik yaddaş (veb: heap)",
                "Universitet",
                "Yüksək",
                "Çox yüksək",
            ],
            [
                "18",
                "Eynizamanlı, asinxron və paralel proqramlaşdırma (veb: paralellik və eyni vaxtda iş)",
                "Universitet",
                "Çox yüksək",
                "Çox yüksək",
            ],
            ["19", "İkilik saylar və rəqəmsal məlumatın təqdimatı", "Hər iki qrup", "Yüksək", "Yüksək"],
            [
                "20",
                "Veb proqramlaşdırma, API və müştəri–server tətbiqləri",
                "Hər iki qrup",
                "Çox yüksək",
                "Yüksək",
            ],
            [
                "21",
                "Əməliyyat sistemləri, proseslər, yaddaş və fayl sistemləri",
                "Universitet",
                "Yüksək",
                "Çox yüksək",
            ],
            [
                "22",
                "Bulud hesablaması, konteynerlər və paylanmış sistemlər",
                "Universitet",
                "Çox yüksək",
                "Çox yüksək",
            ],
            [
                "23",
                "Proqram mühəndisliyi, Git və birgə proqram hazırlama",
                "Yuxarı siniflər və universitet",
                "Yüksək",
                "Yüksək",
            ],
            [
                "24",
                "Kompüter arxitekturası, prosessor, yaddaş və əmrlər",
                "Yuxarı siniflər və universitet",
                "Orta–yüksək",
                "Yüksək",
            ],
            [
                "25",
                "Dinamik proqramlaşdırma və mürəkkəb alqoritmlərin qurulması (veb: irəliləmiş alqoritm quruluşu)",
                "Universitet",
                "Orta–yüksək",
                "Çox yüksək",
            ],
            [
                "26",
                "Kriptoqrafiya, şifrələmə, heşləmə və rəqəmsal imzalar",
                "Yuxarı siniflər və universitet",
                "Çox yüksək",
                "Çox yüksək",
            ],
            [
                "27",
                "Avtomatlar, hesablana bilmə və mürəkkəblik sinifləri",
                "Universitet",
                "Orta",
                "Çox yüksək",
            ],
            [
                "28",
                "Etika, alqoritmik qərəz və hesablamanın sosial təsirləri",
                "Hər iki qrup",
                "Çox yüksək",
                "Yüksək",
            ],
            [
                "29",
                "Oyunların hazırlanması və interaktiv qrafika",
                "Məktəblilər / şagirdlər və yeni başlayanlar",
                "Çox yüksək",
                "Orta–yüksək",
            ],
            [
                "30",
                "Robot texnikası (veb: robototexnika), fiziki hesablama və Əşyaların İnterneti",
                "Məktəblilər / şagirdlər və universitet",
                "Yüksək",
                "Yüksək",
            ],
        ],
    )


def write_school(document: Document) -> None:
    document.add_heading("5. Məktəblilər və orta məktəb şagirdləri üçün", level=1)
    p(
        document,
        "Bu mövzular şagirdlər və orta məktəb şagirdləri üçün xüsusilə faydalıdır. "
        "Məktəb auditoriyası üçün konkret situasiyadan başlamaq və sonra ümumi anlayışa "
        "keçmək əlverişlidir. Yenidən işlənmiş sənəd eyni 23 bəndi altı tematik qrupda "
        "toplayır; veb səhifə isə onları nömrələyib hər bəndə ayrıca nümunə verir. "
        "Aşağıda əvvəlcə tam siyahı, sonra eyni sırada izah və nümunələr gəlir. "
        "Nümunələr materialın qurulmasına kömək edən təkliflərdir, məcburi tapşırıq deyil.",
    )

    document.add_heading("5.1 Mövzu siyahısı", level=2)
    numbered(
        document,
        [
            "Alqoritm nədir — və gündəlik məsələni addımlara necə çevirməli",
            "Bölmə, qanunauyğunluq və abstraksiya",
            "Dəyişənlər — dəyişən saxlanılan dəyərlər",
            "Mənimsətmə ilə riyazi bərabərliyin fərqi",
            "Şərtlər və Bul ifadələri",
            "İç-içə şərtlər və mürəkkəb məntiq",
            "Dövrlər necə işləyir və sonsuz dövr nədən yaranır",
            "Funksiya ilə sadəcə nəticə çap edən əmr",
            "Təsadüfi kod dəyişmək əvəzinə sistemli sazlama",
            "İkilik ədədlər, bit, bayt və onaltılıq yazılış",
            "Mətn, şəkil, səs və video rəqəmsal necə təqdim olunur",
            "İtkili və itkisiz sıxılma",
            "İnternet, Veb və brauzerin fərqi",
            "IP ünvanı, DNS, marşrutlaşdırıcı, paketlər və HTTP",
            "Parollar, fişinq, zərərli proqram, məxfilik və iki amilli autentifikasiya",
            "Şifrələmə, heşləmə və rəqəmsal imzalar",
            "Süni intellekt adi proqramdan nə ilə fərqlənir",
            "Təlim verilənləri və maşın öyrənməsində nümunələrin rolu",
            "Süni intellekt niyə hallüsinasiya edir və ya qərəzi təkrarlayır",
            "Süni intellektin verdiyi məlumatı və kodu necə yoxlamalı",
            "Verilənlərin toplanması, qrafiklər, yanıltıcı təsvir və qərəz",
            "Sayt, oyun, mobil tətbiq və robototexnikanın əsasları",
            "Müəllif hüququ, rəqəmsal kimlik, dezinformasiya və məsuliyyətli istifadə",
        ],
    )
    note_box(
        document,
        "**CSTA 2026.** Yeni CSTA standartları məktəb informatikasını alqoritmlər və "
        "layihələndirmə, proqramlaşdırma, verilənlər və təhlil, sistemlər və təhlükəsizlik, "
        "həmçinin hesablama və cəmiyyət ətrafında qurur. Onlar süni intellekti prioritet "
        "sayır və süni intellekt, kibertəhlükəsizlik, verilənlər elmi, oyun işlənməsi, "
        "fiziki hesablama və proqram təminatı üzrə dərinləşmə istiqamətləri müəyyən edir. "
        "Bunlar ali təhsil ixtisasları deyil, məktəb səviyyəsində ixtisaslaşma istiqamətləridir. "
        "(CSTA 2026 Standards) [1]",
    )

    document.add_heading("5.2 Siyahıdakı mövzuların izahı", level=2)
    p(
        document,
        "Öyrənən hər əmrdən sonra hansı qiymətin dəyişdiyini görə bilməlidir. "
        "Bunun üçün icra addımlarını kiçik cədvəldə izləmək olar. Səhv axtararkən əvvəlcə "
        "gözlənilən nəticə müəyyənləşdirilir, sonra faktiki nəticə ilə müqayisə olunur.",
    )

    document.add_heading("Kompüter kimi düşünmək", level=3)
    document.add_heading("1. Alqoritm nədir — və gündəlik məsələni addımlara necə çevirməli", level=3)
    p(
        document,
        "**Alqoritm** məqsədə çatmaq üçün müəyyən ardıcıllıqla yerinə yetirilən, sonlu və "
        "aydın addımlardır. Çətin olan sözün özü deyil, gündəlik işi maşının təxmin etmədən "
        "yerinə yetirə biləcəyi addımlara çevirməkdir.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Çay dəmləmək: suyu qaynat, stəkana paket qoy, suyu tök, "
        "gözlə, paketi çıxar. Hansısa addım əskikdirsə — «nə qədər gözlə?» — alqoritm natamamdır.",
    )

    document.add_heading("2. Bölmə, qanunauyğunluq və abstraksiya", level=3)
    p(
        document,
        "**Bölmə** böyük məsələni kiçik hissələrə ayırmaqdır. **Qanunauyğunluq** təkrarlananı "
        "görməkdir. **Abstraksiya** həmin məqsəd üçün vacib xüsusiyyətləri saxlayıb ikinci "
        "dərəcəli detalları kənara qoymaqdır.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** Məktəbə ən qısa yolu seçərkən küçələrin əlaqəsi "
        "və məsafəsi vacibdir, binaların rəngi isə adətən vacib deyil.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Məktəb gəzintisini planlamaq: biletlər, yemək və cədvəl "
        "(bölmə); hər avtobusun çıxış vaxtı olmalıdır (qanunauyğunluq); oturacağın rəngini yox, "
        "«neçə yer var» sualını saxlamaq (abstraksiya).",
    )

    document.add_heading("İlk proqramlar", level=3)
    document.add_heading("3. Dəyişənlər — dəyişən saxlanılan dəyərlər", level=3)
    p(
        document,
        "**Dəyişən** proqramın sonradan dəyişə biləcəyi dəyəri saxlayan adlı yerdir. "
        "Bu, cəbrdən gələn sirli simvol deyil, adı olan yaddaşdır.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** `score = 0`, sonra `score = score + 10`. "
        "`score` adlı yerdə artıq 10 var.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd, həm də 3.6-cı bölmə).** `a = 5`, `b = a`, `a = 8`. "
        "Sonda `b` 5 qalır, çünki `b`-yə `a`-nın həmin andakı qiyməti yazılıb.",
    )

    document.add_heading("4. Mənimsətmə ilə riyazi bərabərliyin fərqi", level=3)
    p(
        document,
        "**Mənimsətmə** dəyəri dəyişənə yazır. Bir çox dildə bərabərlik işarəsi riyaziyyatdakı "
        "kimi «hər iki tərəf artıq eynidir» demir.",
    )
    example_box(
        document,
        "**Nümunə (hər iki mənbə).** `x = x + 1` tənlik kimi mənasızdır, mənimsətmə kimi isə "
        "«indiki `x`-i götür, bir əlavə et, nəticəni yenə `x`-ə yaz» deməkdir.",
    )

    document.add_heading("5. Şərtlər və Bul ifadələri", level=3)
    p(
        document,
        "**Şərt** bəli-xeyr sualı verir. **Bul** (*Boolean*) dəyəri yalnız doğrudur və ya "
        "yanlışdır. Proqram növbəti addımı bu cavaba görə seçir.",
    )
    example_box(
        document,
        "**Nümunə.** Əgər temperatur > 30-dursa, «isti» yaz; əks halda «normal» yaz.",
    )

    document.add_heading("6. İç-içə şərtlər və mürəkkəb məntiq", level=3)
    p(
        document,
        "Şərtlər bir-birinin içində ola bilər, bir neçə yoxlama isə «və» və ya «və ya» ilə "
        "birləşdirilə bilər. Çətinlik qərar sırasını aydın saxlamaqdır.",
    )
    example_box(
        document,
        "**Nümunə.** Həftə içidirsə və yağış yağırsa, avtobusa min; həftə içidirsə və quru "
        "havadırsa, piyada get; əks halda evdə qal.",
    )

    document.add_heading("7. Dövrlər necə işləyir və sonsuz dövr nədən yaranır", level=3)
    p(
        document,
        "**Dövr** eyni addımları təkrarlayır. Onun dayanma yolu olmalıdır. "
        "**Sonsuz dövr** dayanma şərti heç vaxt doğru olmayanda baş verir.",
    )
    example_box(
        document,
        "**Nümunə.** «Çən dolmayana qədər su tök» yalnız hər tökülən su səviyyəni qaldırırsa "
        "işləyir. Kran bağlıdırsa, dövr bitmir.",
    )

    document.add_heading("8. Funksiya ilə sadəcə nəticə çap edən əmr", level=3)
    p(
        document,
        "**Funksiya** girişlər (**parametrlər**) ala və nəticə (**qaytarılan qiymət**) verə "
        "bilən adlı iş parçasıdır. Ekrana rəqəm yazmaq onu sonrakı istifadə üçün qaytarmaq deyil.",
    )
    example_box(
        document,
        "**Nümunə.** `add(2, 3)` funksiyası təkcə «5» göstərməməli, 5-i qaytarmalıdır ki, "
        "proqramın başqa hissəsi onu saxlaya bilsin. Biri yalnız ekrana yazan, digəri isə "
        "qiymət qaytaran iki qısa funksiya bu fərqi görünən edir.",
    )

    document.add_heading("9. Təsadüfi kod dəyişmək əvəzinə sistemli sazlama", level=3)
    p(
        document,
        "**Sazlama** (*debugging*) səhvləri təsadüfi sətir dəyişməklə yox, fərziyyəni "
        "yoxlamaqla tapmaqdır. Əvvəlcə gözlənilən nəticə müəyyənləşdirilir, sonra faktiki "
        "nəticə ilə müqayisə olunur. Bunu öyrənmək də, öyrətmək də çətindir, çünki təcrübəli "
        "proqramçılar çox vaxt açıqlamadıqları üsullara arxalanır. [10]",
    )
    example_box(
        document,
        "**Nümunə.** Cəm səhvdirsə, əvvəl birinci girişi, sonra toplamanı, sonra göstərişi — "
        "bir-bir — yoxla; bütün proqramı yenidən yazma.",
    )

    document.add_heading("Kompüterlər məlumatı necə saxlayır", level=3)
    document.add_heading("10. İkilik ədədlər, bit, bayt və onaltılıq yazılış", level=3)
    p(
        document,
        "Kompüterlər məlumatı **bit**lərlə — 0 və ya 1 — saxlayır. Səkkiz bit bir **bayt**dır. "
        "**Onaltılıq** (16-lıq) eyni ikilik dəyərləri qısa yazmaq üsuludur.",
    )
    example_box(
        document,
        "**Nümunə.** Adi mətn kodlaşdırmasında A hərfi 01000001 baytıdır, onaltılıqda çox vaxt 41 yazılır.",
    )

    document.add_heading("11. Mətn, şəkil, səs və video rəqəmsal necə təqdim olunur", level=3)
    p(
        document,
        "Kompüter şəkli «görmür», mahnını «eşitmir». O, rəng, parlaqlıq və ya hava təzyiqinin "
        "nümunələrini ədədlərlə saxlayır, sonra həmin ədədlərdən media yenidən qurulur. "
        "Təsviri piksellər və onların rəng qiymətləri ilə izah etmək mümkündür.",
    )
    example_box(
        document,
        "**Nümunə.** Ağ-qara şəkil 0 və 1-lərdən ibarət tor ola bilər: 0 ağ, 1 qara.",
    )

    document.add_heading("12. İtkili və itkisiz sıxılma", level=3)
    p(
        document,
        "**Sıxılma** faylı kiçildir. **Itkisiz** sıxılmada ilkin məlumat tam bərpa olunur. "
        "**İtkili** sıxılmada həcmi azaltmaq üçün məlumatın bir hissəsi atılır.",
    )
    example_box(
        document,
        "**Nümunə.** Mətnin ZIP arxivi adətən itkisizdir. Güclü sıxılmış foto bir qədər yumşaq "
        "görünə bilər, çünki rəng təfərrüatının bir hissəsi atılıb.",
    )

    document.add_heading("Şəbəkələr və təhlükəsizlik", level=3)
    document.add_heading("13. İnternet, Veb və brauzerin fərqi", level=3)
    p(
        document,
        "**İnternet** bir-biri ilə əlaqələnmiş şəbəkələr sistemidir. **Veb** həmin "
        "infrastruktur üzərində işləyən xidmətlərdən biridir: ünvanlarla bağlı səhifələr. "
        "**Brauzer** isə veb məzmununa baxmağa imkan verən proqramdır.",
    )
    example_box(
        document,
        "**Nümunə.** İnternet yol şəbəkəsidir; Veb yol kənarındakı mağazalardır; brauzer "
        "mağazaya getdiyin avtomobildir.",
    )

    document.add_heading("14. IP ünvanı, DNS, marşrutlaşdırıcı, paketlər və HTTP", level=3)
    p(
        document,
        "**IP ünvanı** maşının şəbəkə nömrəsidir. **DNS** (*Domain Name System*) domen adını "
        "şəbəkə ünvanına uyğunlaşdırır. **Marşrutlaşdırıcı** **paketləri** — kiçik məlumat "
        "dilimlərini — ötürür. **HTTP** (*Hypertext Transfer Protocol*) veb resursları üçün "
        "sorğu və cavab qaydalarını müəyyənləşdirir.",
    )
    example_box(
        document,
        "**Nümunə.** Səhifə ünvanını yazanda əvvəl DNS-dən nömrə soruşulur, sonra HTTP həmin "
        "maşından səhifəni istəyir, səhifə paketlər şəklində gəlir.",
    )

    document.add_heading(
        "15. Parollar, fişinq, zərərli proqram, məxfilik və iki amilli autentifikasiya",
        level=3,
    )
    p(
        document,
        "**Fişinq** aldadıcı mesaj və ya saytla parol və ya başqa sirri əldə etməyə cəhddir. "
        "**Zərərli proqram** zərər vurmaq və ya casusluq etmək üçün yazılmış proqramdır. "
        "**İki amilli autentifikasiya** (veb səhifədə: iki faktorlu təsdiq) fərqli "
        "kateqoriyalardan iki yoxlama vasitəsindən istifadə edir.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** Məktəb hesabına girişdə parol birinci yoxlamadır; "
        "ayrıca cihaz və ya tətbiq vasitəsilə təsdiq ikinci amil ola bilər.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Məktəb portalına bənzəyən, amma naməlum saytda «parolunu "
        "təsdiqlə» deyən mesaj adi fişinq nümunəsidir.",
    )

    document.add_heading("16. Şifrələmə, heşləmə və rəqəmsal imzalar", level=3)
    p(
        document,
        "**Şifrələmə** məlumatın məxfiliyini qoruyur: onu yalnız düzgün açarı olan oxuya bilir. "
        "**Heşləmə** məlumatın qısa rəqəmsal izini yaradır; izdən əsli faydalı şəkildə bərpa "
        "etmək olmur. **Rəqəmsal imza** bütövlüyü və mənşəyi yoxlamağa xidmət edir.",
    )
    example_box(
        document,
        "**Nümunə.** Mesajlaşma tətbiqi mətni yolda şifrələyə bilər. Sayt parolun özünü yox, "
        "heşini saxlaya bilər.",
    )

    document.add_heading("Süni intellekt və verilənlər", level=3)
    document.add_heading("17. Süni intellekt adi proqramdan nə ilə fərqlənir", level=3)
    p(
        document,
        "Adi proqram insanın yazdığı qaydaları izləyir. Bir çox **süni intellekt** sistemi, "
        "xüsusən **maşın öyrənməsi**, davranışını tam yazılı qayda kitabından yox, "
        "**təlim verilənlərindən** — böyük nümunə dəstlərindən — qurur. Süni intellektin "
        "bütün üsulları eyni qaydada işləmir; izahda konkret olaraq maşın öyrənməsinə "
        "əsaslanan sistemləri götürmək daha aydındır.",
    )
    example_box(
        document,
        "**Nümunə.** Spam süzgəcinə bütün mümkün zibil cümlələr verilmir. Ona çoxlu spam və "
        "adi məktub nümunəsi göstərilir, sonra yeni məktubun hansına daha çox bənzədiyi "
        "qiymətləndirilir.",
    )

    document.add_heading("18. Təlim verilənləri və maşın öyrənməsində nümunələrin rolu", level=3)
    p(
        document,
        "Təlimdə istifadə olunan nümunələr sistemin nə edə biləcəyini və harada yanılacağını "
        "formalaşdırır. Nümunələr dar və ya qərəzli olsa, cavablar da belə olacaq.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Yalnız yay fotoları ilə öyrədilmiş model qarı şəkildəki "
        "qüsur kimi səhv adlandıra bilər.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** Yalnız bir şəraitdə çəkilmiş şəkillərlə "
        "öyrədilmiş model başqa şəraitdə zəif nəticə göstərə bilər.",
    )

    document.add_heading("19. Süni intellekt niyə hallüsinasiya edir və ya qərəzi təkrarlayır", level=3)
    p(
        document,
        "**Hallüsinasiya** süni intellekt sisteminin inandırıcı görünən, lakin yanlış və ya "
        "əsassız məlumat yaratmasıdır. Sistem yoxlanılmış fakta baxmır, ehtimal olunan söz "
        "sırasını proqnozlaşdırır. Təlim verilənlərindəki **qərəzi də təkrarlaya** bilər. "
        "İzahda düzgün səslənən cavabla yoxlanmış cavab arasındakı fərq göstərilməlidir.",
    )
    example_box(
        document,
        "**Nümunə.** Söhbət botu real görünən məqalə adı uydura bilər, və ya verilənlərdə "
        "tez-tez görünən stereotipi təkrarlaya bilər.",
    )

    document.add_heading("20. Süni intellektin verdiyi məlumatı və kodu necə yoxlamalı", level=3)
    p(
        document,
        "Süni intellektin çıxışını qaralama sayın. İddiaları etibarlı mənbə ilə tutuşdurun, "
        "koda isə etibar etməzdən əvvəl onu işə salın. Mənbə bunu yalnız mütəxəssis yox, "
        "məktəb səviyyəli bacarıq kimi göstərir. Stack Overflow 2025 nəticələri də bu "
        "yoxlamanı praktik mövzu kimi əsaslandırır. [4, 26]",
    )
    example_box(
        document,
        "**Nümunə.** Model tarix və ya düstur yazırsa, istinad olunan səhifəni açın və ya "
        "düsturu tanış nümunə ilə yoxlayın.",
    )

    document.add_heading("21. Verilənlərin toplanması, qrafiklər, yanıltıcı təsvir və qərəz", level=3)
    p(
        document,
        "Süni intellektsiz də verilənlər yanıltmağa bilər. Qrafik miqyası gizlədə, seçmə bir "
        "qrupu buraxa, rəngli diaqram isə zəif qanunauyğunluğu güclü göstərə bilər.",
    )
    example_box(
        document,
        "**Nümunə.** 0-dan yox, 90-dan başlayan sütun diaqramı yaxın iki nəticəni uzaq göstərə bilər.",
    )

    document.add_heading("Nəsə qurmaq", level=3)
    document.add_heading("22. Sayt, oyun, mobil tətbiq və robototexnikanın əsasları", level=3)
    p(
        document,
        "Bunlar yuxarıdakı ideyaların real auditoriya ilə qarşılaşdığı ilk praktik mühitlərdir: "
        "açılan səhifə, xal toplayan oyun, ayarı saxlayan telefon tətbiqi, hiss edib hərəkət "
        "edən robot. Mənbə onları əsaslar kimi göstərir; hər şagirdin dördünü də qurması tələb "
        "deyil. Sadə oyun layihəsi şərtləri, dövrləri və hadisələri birlikdə izah edə bilər. "
        "Robot və sensor layihələri proqramın fiziki mühitdə yaratdığı nəticəni görünən edir.",
    )
    example_box(
        document,
        "**Nümunə.** Sadə oyunda belə dəyişən (xal), dövr (oyun davam edir) və şərt (udulub) lazımdır.",
    )

    document.add_heading("Hesablama və cəmiyyət", level=3)
    document.add_heading(
        "23. Müəllif hüququ, rəqəmsal kimlik, dezinformasiya və məsuliyyətli istifadə",
        level=3,
    )
    p(
        document,
        "Hesablama yalnız texnika deyil. Şagirdlər əsərin kimə məxsus olduğunu, onlayn "
        "kimliyin necə kopyalana biləcəyini, yalan iddiaların necə yayıldığını və sinifdə "
        "və ya açıq paylaşımda məsuliyyətli istifadənin nə olduğunu da bilməlidirlər. "
        "Layihənin şəkil, səs və mətn mənbələrinin göstərilməsi müəllif hüquqları ilə "
        "texniki işi əlaqələndirir.",
    )
    example_box(
        document,
        "**Nümunə.** Şəkli məktəb işinə kredit vermədən qoymaq, onu yükləmək asan olsa belə, "
        "müəllif hüququnu poza bilər.",
    )


def write_university(document: Document) -> None:
    document.add_heading("6. Universitet tələbələri üçün", level=1)
    p(
        document,
        "Bu auditoriya üçün təkcə mexanizmin işləməsini deyil, alternativ həllər arasındakı "
        "fərqi və seçim səbəblərini də izah etmək vacibdir. Veb səhifə mövzuları kollec və "
        "universitet təhsili üçün xüsusilə vacib sayır; yenidən işlənmiş sənəd eyni bəndləri "
        "yeddi qrupda birləşdirir.",
    )
    note_box(
        document,
        "**Quruluş fərqi.** Veb səhifə 25 bəndi belə düzür: icra (1–5), verilənlər "
        "strukturları (6–9), bazalar (10–12), sistemlər (13–18), maşın öyrənməsi (19–21), "
        "proqram işi (22–24), nəzəriyyə (25). Yenidən işlənmiş sənəd paradiqmaları və "
        "kompilyatorları icra qrupuna, avtomatlar, P və NP-ni isə alqoritmlər qrupuna salır. "
        "Məzmun eynidir; aşağıdakı nömrələmə veb səhifədəki 25 bəndi saxlayır ki, siyahı "
        "izahlarla eyni sırada qalsın.",
    )

    document.add_heading("6.1 Mövzu siyahısı", level=2)
    numbered(
        document,
        [
            "Proqramın icra modelləri və vəziyyətin izlənməsi",
            "Əhatə dairəsi, ömür, parametr ötürmə və təxəllüs",
            "İstinadlar, göstəricilər və yaddaş idarəsi",
            "Rekursiya və rekursiv çağırış steki",
            "Obyekt kimliyi, siniflər, obyektlər, irsiyyət və polimorfizm",
            "Siyahılar, bağlı strukturlar, steklər, növbələr, ağaclar, heap-lər və qraflar",
            "Sıralama, axtarış və qraf alqoritmləri",
            "Big-O, ən pis hal və yaddaş–vaxt mübadiləsi",
            "Acgöz alqoritmlər, böl və həll et, dinamik proqramlaşdırma",
            "Relyasion modelləşdirmə, açarlar, birləşmələr və normallaşdırma",
            "SQL-də aqreqasiya, alt sorğular və pəncərə funksiyaları",
            "Tranzaksiyalar, izolyasiya, kilidləmə və verilənlər bazasının uyğunluğu",
            "Proseslər, axınlar, planlaşdırma və virtual yaddaş",
            "Yarış vəziyyətləri, sinxronlaşdırma, qarşılıqlı kilidlənmə və qeyri-determinizm",
            "Şəbəkə protokolları, marşrutlaşdırma, etibarlılıq və sıxlıq",
            "Autentifikasiya, avtorizasiya, şifrələmə və təhlükəsiz proqram layihələndirməsi",
            "Paylanmış sistemlər, replikasiya, uyğunluq və konsensus",
            "Bulud hesablaması, konteynerlər, orkestrasiya və yerləşdirmə",
            "Maşın öyrənməsinin qiymətləndirilməsi, həddindən artıq və yetərsiz öyrənmə",
            "Verilənlər sızması, sinif balanssızlığı, model qərəzi və izah edilə bilmə",
            "Böyük dil modellərində tokenləşdirmə, vektor təmsilləri, diqqət, axtarış və hallüsinasiya",
            "Proqramın yoxlanması, arxitektura, versiyaların idarə edilməsi və davamlı inteqrasiya",
            "Funksional, məntiqi və eynizamanlı proqramlaşdırma paradiqmaları",
            "Kompilyatorlar, interpretatorlar və proqramlaşdırma dillərinin semantikası",
            "Avtomatlar, hesablana bilmə, reduksiya, P, NP və NP-tamlıq",
        ],
    )

    document.add_heading("6.2 Siyahıdakı mövzuların izahı", level=2)

    document.add_heading("Proqramlar necə işləyir", level=3)
    document.add_heading("1. Proqramın icra modelləri və vəziyyətin izlənməsi", level=3)
    p(
        document,
        "Proqram **vəziyyətin** — dəyişənlərin cari qiymətlərinin və növbəti əmrin — "
        "ardıcıl dəyişməsidir. **İzləmə** hər addımdan sonra həmin vəziyyəti yazmaqdır. "
        "Sonrakı bir çox yanlış anlama bu prosesin zehni modelinin olmaması ilə bağlanır. [8, 9]",
    )
    example_box(
        document,
        "**Nümunə.** `x = 2`; `x = x * 3` üçün iz: əvvəl `x` yoxdur; birinci sətirdən sonra "
        "`x` 2-dir; ikinci sətirdən sonra `x` 6-dır.",
    )

    document.add_heading("2. Əhatə dairəsi, ömür, parametr ötürmə və təxəllüs", level=3)
    p(
        document,
        "**Əhatə dairəsi** və ya **görünmə sahəsi** adın harada əlçatan olduğunu bildirir. "
        "**Ömür** (*yaşam müddəti*) saxlanan dəyərin və ya obyektin nə qədər mövcud "
        "olduğudur. **Parametr ötürmə** girişlərin funksiyaya necə düşməsidir. "
        "**Təxəllüs** (*aliasing*) iki adın eyni obyektə işarə etməsidir; ona görə bir adla "
        "edilən dəyişiklik o biri addan da görünür.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Funksiya siyahı alıb element əlavə edirsə, çağıranın "
        "siyahısı da dəyişə bilər — hər iki ad eyni siyahıya baxırdı.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** İki dəyişən eyni obyektə yönəlirsə, obyekt "
        "dəyişdikdə hər iki istinad vasitəsilə həmin dəyişiklik görünür. İstinadların birini "
        "başqa obyektə yönəltmək isə fərqli əməliyyatdır.",
    )

    document.add_heading("3. İstinadlar, göstəricilər və yaddaş idarəsi", level=3)
    p(
        document,
        "**İstinad** və ya **göstərici** dəyərin özünü yox, yerini saxlayır; dəqiq davranış "
        "dildən asılıdır. **Yaddaş idarəsi** həmin yerin nə vaxt yaradılıb azad olunacağına "
        "qərardır. Dolayı müraciət, eyni obyektə bir neçə istinad, görünmə sahəsi və parametr "
        "ötürülməsi uzunmüddətli çətinlik yarada bilər. [14]",
    )
    example_box(
        document,
        "**Nümunə.** İki dəyişən eyni müştəri qeydinə baxa bilər. Ünvanı hansı dəyişənlə "
        "dəyişsən, eyni qeyd dəyişir.",
    )

    document.add_heading("4. Rekursiya və rekursiv çağırış steki", level=3)
    p(
        document,
        "**Rekursiya** funksiyanın birbaşa və ya dolayı yolla özünü çağırmasıdır; əlavə "
        "çağırış tələb etməyən **baza halı**na qədər. Hər çağırış daxili çağırış bitənə qədər "
        "**çağırış steki**ndə gözləyir. İzahda həm yeni çağırışların yaranmasını, həm də "
        "nəticələrin geri qayıtmasını göstərmək lazımdır. [13]",
    )
    example_box(
        document,
        "**Nümunə.** `factorial(4)` `factorial(3)`-ü, o da `factorial(2)`-ni gözləyir, "
        "ta `factorial(1)` 1 qaytarana qədər.",
    )

    document.add_heading("5. Obyekt kimliyi, siniflər, obyektlər, irsiyyət və polimorfizm", level=3)
    p(
        document,
        "**Sinif** təsvirdir; **obyekt** isə bir nüsxədir. **Kimlik** iki "
        "obyektin eyni ədədləri saxlayıb-saxlamamasından çox, hansı obyekt olduğunu soruşur. "
        "**İrsiyyət** bir sinfin o birinin quruluşunu yenidən istifadəsinə imkan verir. "
        "**Polimorfizm** (veb səhifədə: çoxformalıq) eyni interfeysin müxtəlif obyektlərdə "
        "fərqli davranış verməsinə imkan yaradır. Konstruktor obyekt yaradılarkən onun "
        "başlanğıc vəziyyətinin qurulmasında iştirak edir. [12]",
    )
    example_box(
        document,
        "**Nümunə.** Dairə və düzbucaqlı hər ikisi «sahə» deyə bilər, amma hər biri onu "
        "fərqli hesablayır.",
    )


    document.add_heading("Verilənlər strukturları və alqoritmlər", level=3)
    document.add_heading(
        "6. Siyahılar, bağlı strukturlar, steklər, növbələr, ağaclar, heap-lər və qraflar",
        level=3,
    )
    p(
        document,
        "**Verilənlər strukturu** bəzi əməliyyatların ucuz qalması üçün dəyərləri təşkil "
        "etmə üsuludur. Siyahı sıranı saxlayır; stek son gələn birinci çıxır; növbə birinci "
        "gələn birinci çıxır; ağac və heap iyerarxiyanı, qraf isə əlaqələri modelləşdirir. "
        "Burada **heap** verilənlər strukturu kimi yığını bildirir; eyni söz yaddaş bölgəsi "
        "üçün də işlənir. Yenidən işlənmiş sənəd bu ikimənalılığı «dinamik yaddaş» ifadəsi ilə "
        "ayırır. [11]",
    )
    example_box(
        document,
        "**Nümunə.** Brauzerin Geri düyməsi stek kimidir: açdığın son səhifəyə birinci qayıdırsan.",
    )

    document.add_heading("7. Sıralama, axtarış və qraf alqoritmləri", level=3)
    p(
        document,
        "Bunlar elementləri sıraya salmaq, tapmaq və ya əlaqələr üzrə gəzmək üçün standart "
        "üsullardır. Tələbə həm üsulu, həm də onun nə vaxt doğru olduğunu görməlidir.",
    )
    example_box(
        document,
        "**Nümunə.** Sıralanmış sinif siyahısında ad tapmaq üçün ikili axtarışdan istifadə "
        "oluna bilər: qalan siyahını dəfələrlə yarıya bölmək.",
    )

    document.add_heading("8. Big-O, ən pis hal və yaddaş–vaxt mübadiləsi", level=3)
    p(
        document,
        "**Big-O** (veb səhifədə: Böyük-O) giriş böyüdükcə alqoritmin resurs tələbinin "
        "artımına yuxarı sərhəd vermək üçün istifadə olunur; dəqiq saniyə sayı deyil. "
        "**Ən pis hal** ən bahalı qanuni girişi soruşur. **Yaddaş–vaxt mübadiləsi** vaxt "
        "qazanmaq üçün daha çox yaddaş, və ya əksi, istifadə etməkdir.",
    )
    example_box(
        document,
        "**Nümunə.** Hər elementi bir dəfə baxmaq xətti vaxtdır. Hər cütə baxmaq siyahı "
        "uzandıqca daha sürətlə bahalaşır.",
    )

    document.add_heading("9. Acgöz alqoritmlər, böl və həll et, dinamik proqramlaşdırma", level=3)
    p(
        document,
        "**Acgöz** üsul yerli ən yaxşı addımı götürür. **Böl və həll et** (veb səhifədə: "
        "böl və hökm et) məsələni parçalayır, hissələri həll edir və birləşdirir. "
        "**Dinamik proqramlaşdırma** kəsişən altməsələlərin nəticələrini saxlayır ki, "
        "yenidən hesablanmasın. İzahda əvvəlcə təkrarlanan işi göstərmək, sonra yadda "
        "saxlamağın onu necə azaltdığını nümayiş etdirmək olar. Üsul seçimi və rekurrent "
        "əlaqə adi çətinliklər sırasındadır. [20, 27]",
    )
    example_box(
        document,
        "**Nümunə.** Ən az sikkə ilə xırdalamaq, sikkə sistemi imkan verirsə, acgöz ola "
        "bilər; bəzi sistemlər daha diqqətli, saxlanılan həll üsulu tələb edir.",
    )

    document.add_heading("Verilənlər və verilənlər bazaları", level=3)
    document.add_heading("10. Relyasion modelləşdirmə, açarlar, birləşmələr və normallaşdırma", level=3)
    p(
        document,
        "**Relyasion** (veb səhifədə: relasion) baza məlumatı cədvəllərdə saxlayır. "
        "**Açar** sətri unikal tanıyır. **Birləşmə** (*JOIN*) bir neçə cədvəlin sətirlərini "
        "birləşdirir. **Normallaşdırma** təkrarlanma və yenilənmə problemlərini azaltmaq "
        "üçün cədvəllərin təşkilidir.",
    )
    example_box(
        document,
        "**Nümunə (hər iki mənbə).** Tələbələr və fənlər cədvəlləri qeydiyyat cədvəlində "
        "görüşür. Eyni tələbənin bir neçə kursa yazılması nəticədə niyə bir neçə sətir "
        "yarandığını göstərir; hər tələbə sətrində fənn adını təkrarlamaq olmaz.",
    )

    document.add_heading("11. SQL-də aqreqasiya, alt sorğular və pəncərə funksiyaları", level=3)
    p(
        document,
        "**SQL** (*Structured Query Language*) verilənlər bazasına sorğu vermək üçün "
        "istifadə olunan dildir. **Aqreqasiya** (veb səhifədə: cəmləmə) çoxlu qiymətdən "
        "cəm, say və orta kimi ümumi nəticə çıxarır. **Alt sorğu** başqa sorğunun içində "
        "yerləşir. **Pəncərə funksiyası** əlaqəli sətirlər üzrə hesab apararaq ayrı "
        "sətirləri nəticədə saxlayır. Genişmiqyaslı təhlildə tələbə səhvləri tez-tez "
        "birləşmə, alt sorğu və GROUP BY ilə bağlıdır. GROUP BY sətirləri ümumi xüsusiyyətə "
        "görə qruplaşdırır; nəticə cədvəlinin necə yarandığını addımlarla göstərmək "
        "faydalıdır. [16]",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd; veb səhifədə bu bəndin ayrıca nümunəsi yoxdur).** "
        "Universitetin tələbə və qeydiyyat cədvəlləri JOIN əməliyyatını izah etməyə yaraya "
        "bilər. Əvvəlcə birləşmənin sətir sayını artırdığını, sonra GROUP BY-nin həmin "
        "sətirləri kurs və ya tələbə üzrə yığdığını göstərmək olar.",
    )

    document.add_heading(
        "12. Tranzaksiyalar, izolyasiya, kilidləmə və verilənlər bazasının uyğunluğu",
        level=3,
    )
    p(
        document,
        "**Tranzaksiya** əlaqəli əməliyyatları vahid iş kimi idarə edir: birlikdə başa "
        "çatmalı və ya birlikdə uğursuz olmalıdır. **İzolyasiya** (veb səhifədə: izolə) "
        "paralel tranzaksiyaların bir-birinin aralıq nəticələrini necə görə bildiyini "
        "müəyyənləşdirir. **Kilidləmə** eyni məlumat üzərində toqquşmanın qarşısını alır. "
        "**Uyğunluq** (veb səhifədə: tutarlılıq) dəyişiklikdən sonra saxlanan faktların "
        "qaydalara əməl etməsidir.",
    )
    example_box(
        document,
        "**Nümunə.** Bank köçürməsi bir hesabdan çıxarış etməməli, o biri hesaba da əlavə "
        "etməyibsə.",
    )

    document.add_heading("Sistemlər", level=3)
    document.add_heading("13. Proseslər, axınlar, planlaşdırma və virtual yaddaş", level=3)
    p(
        document,
        "**Proses** icra olunan proqramın mühitidir; adətən öz yaddaşı olur. **Axın** "
        "həmin mühitdə icra yoludur. **Planlaşdırma** prosessorun növbəti işinin hansı "
        "olacağına qərar verir. **Virtual yaddaş** proqramın gördüyü ünvanları fiziki "
        "yaddaşdan ayırır və proqramın fiziki RAM-dən böyük və ya ondan təcrid olunmuş "
        "ünvan sahəsi istifadə etməsinə imkan verir. Tələbələr bu mexanizmləri əzbərləyə, "
        "amma həllər arasında seçim etməkdə və mübadiləni izah etməkdə çətinlik çəkə "
        "bilərlər. [17]",
    )
    example_box(
        document,
        "**Nümunə.** Brauzer hər vərəqə üçün ayrı proses saxlaya bilər ki, bir səhifənin "
        "çökməsi o birilərini aparmasın.",
    )

    document.add_heading(
        "14. Yarış vəziyyətləri, sinxronlaşdırma, qarşılıqlı kilidlənmə və qeyri-determinizm",
        level=3,
    )
    p(
        document,
        "Bir neçə işin icrası üst-üstə düşdükdə nəticə addımların ardıcıllığından asılı "
        "ola bilər. **Yarış vəziyyəti** (veb səhifədə: yarış şərti) nəticənin icra "
        "ardıcıllığından asılı olmasıdır. **Sinxronlaşdırma** axınları əlaqələndirir. "
        "**Qarşılıqlı kilidlənmə** (veb səhifədə: tıxac) işlərin bir-birinin saxladığı "
        "resursları gözləyərək dayanmasıdır. **Qeyri-determinizm** səhv proqramın bəzi "
        "işəsalmalarda düzgün görünə bilməsi deməkdir. Proqramın bir dəfə düzgün işləməsi "
        "bütün mümkün icra ardıcıllıqlarında düzgün olduğunu sübut etmir. [15]",
    )
    note_box(
        document,
        "**Fərq.** Veb səhifə deadlock üçün «tıxac», şəbəkə yüklənməsi üçün isə «tıxaclıq» "
        "yazır. Yenidən işlənmiş sənəd birincini «qarşılıqlı kilidlənmə», ikincisini "
        "«sıxlıq» adlandırır. Eyni «tıxac» kökü iki fərqli anlayışı qarışdıra bilər.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** İki axın da «1 yer qalıb» oxuyur və hər ikisi son yeri satır.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** İki axının eyni sayğacı yeniləməsini addımlara "
        "bölmək yarış vəziyyətini göstərir.",
    )

    document.add_heading("15. Şəbəkə protokolları, marşrutlaşdırma, etibarlılıq və sıxlıq", level=3)
    p(
        document,
        "**Protokol** maşınlar arasında razılaşdırılmış söhbətdir. **Marşrutlaşdırma** yolu "
        "seçir. **Etibarlılıq** itən məlumatın tapılıb-tapılmayacağını və yenidən "
        "göndəriləcəyini soruşur. **Sıxlıq** (veb səhifədə: tıxaclıq) ortaq yolun həddən "
        "artıq yüklənməsidir.",
    )
    example_box(
        document,
        "**Nümunə.** Video zəng yol tıxandıqda əbədi donmaq əvəzinə şəkil keyfiyyətini sala bilər.",
    )

    document.add_heading(
        "16. Autentifikasiya, avtorizasiya, şifrələmə və təhlükəsiz proqram layihələndirməsi",
        level=3,
    )
    p(
        document,
        "**Autentifikasiya** «Siz kimsiniz?», **avtorizasiya** isə «Nə etməyə icazəniz var?» "
        "sualına cavab verir. Şifrələmə məlumatı yolda və ya saxlancda qoruyur. Təhlükəsiz "
        "layihə bunları əvvəldən sistemin hissəsi sayır, sonda bəzək kimi yox. "
        "Kibertəhlükəsizlikdə zəiflik, təhdid, risk və qorunma tədbiri qarışdırıla bilər. "
        "Zəiflik sistemdəki nöqsan, təhdid zərər yarada bilən amil, risk isə mümkün zərərin "
        "ehtimalı və nəticələri ilə bağlı anlayışdır. [19]",
    )
    example_box(
        document,
        "**Nümunə.** Giriş kimliyi sübut edir; tələbə olmaq başqa tələbənin qiymətini dəyişmək "
        "hüququ vermir. Sistemə daxil olmuş istifadəçinin bütün faylları dəyişmək hüququ "
        "olmaya bilər.",
    )

    document.add_heading("17. Paylanmış sistemlər, replikasiya, uyğunluq və konsensus", level=3)
    p(
        document,
        "Paylanmış sistemdə hissələr müxtəlif kompüterlərdə işləyir və şəbəkə ilə əlaqə "
        "saxlayır. **Replikasiya** məlumatın surətlərini müxtəlif yerlərdə saxlamaqdır. "
        "Nüsxələr razılaşmaya bilər, ona görə **uyğunluq** (veb səhifədə: tutarlılıq) "
        "qaydası və bəzən növbəti dəyər üçün **konsensus** (veb səhifədə: razılaşma) üsulu "
        "lazımdır.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** İki surətin müvəqqəti əlaqəsiz qalması "
        "məlumatların uyğunlaşdırılmasını izah etmək üçün yaxşı başlanğıcdır.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Bir neçə serverdə saxlanan sənəddə iki nəfər eyni anda "
        "yazırsa, kimin düzəlişinin qalacağına qərar verilməlidir.",
    )

    document.add_heading("18. Bulud hesablaması, konteynerlər, orkestrasiya və yerləşdirmə", level=3)
    p(
        document,
        "**Bulud hesablaması** server, saxlanc və şəbəkəni xidmət kimi icarəyə verir. "
        "**Konteyner** tətbiqi asılılıqları ilə birlikdə işlətməyi asanlaşdırır. "
        "**Orkestrasiya** (veb səhifədə: orkestrləşdirmə) çoxlu konteynerin yerləşdirilməsi, "
        "işə salınması və idarə edilməsinin əlaqələndirilməsidir. **Yerləşdirmə** yeni "
        "versiyanı xidmətə qoymaqdır.",
    )
    example_box(
        document,
        "**Nümunə.** Kurs saytı konteynerdə işləyə bilər ki, eyni obraz həm noutbukda, həm "
        "də buludda işləsin.",
    )

    document.add_heading("Maşın öyrənməsi və iri modellər", level=3)
    document.add_heading(
        "19. Maşın öyrənməsinin qiymətləndirilməsi, həddindən artıq və yetərsiz öyrənmə",
        level=3,
    )
    p(
        document,
        "**Həddindən artıq öyrənmə** (*overfitting*) modelin təlim məlumatının özəlliklərinə "
        "həddən çox uyğunlaşıb yeni nümunələrdə yanılmasıdır. **Yetərsiz öyrənmə** "
        "(*underfitting*) modelin əsas əlaqələri kifayət qədər öyrənməməsidir. Qiymətləndirmə "
        "modelin görmədiyi verilənlərlə aparılmalıdır.",
    )
    example_box(
        document,
        "**Nümunə.** Keçən ilin sinifindəki hər tələbəni tanıyan, amma yeni tələbəni "
        "yerləşdirə bilməyən model artıq uyğunlaşıb.",
    )

    document.add_heading(
        "20. Verilənlər sızması, sinif balanssızlığı, model qərəzi və izah edilə bilmə",
        level=3,
    )
    p(
        document,
        "**Verilənlər sızması** modelə proqnoz anında olmayacağı məlumatı təsadüfən verməkdir; "
        "məsələn, sınaq məlumatının təlim prosesinə qarışması nəticələri olduğundan yaxşı "
        "göstərə bilər. **Sinif balanssızlığı** nümunə kateqoriyalarının sayca qeyri-bərabər "
        "olmasıdır. **Qərəz** nəticələrin müəyyən qruplar və ya şərait üçün sistemli şəkildə "
        "təhrif olunması ola bilər. **İzah edilə bilmə** (veb səhifədə: izahlılıq) modelin "
        "nəticəsinə təsir edən amillərin insanlar üçün başadüşülən şəkildə açıqlana "
        "bilməsidir.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd).** Sınaq məlumatının təlimə qarışması "
        "qiymətləndirmə kontekstində verilənlər sızmasına nümunədir.",
    )
    example_box(
        document,
        "**Nümunə (veb səhifə).** Tibbi model xəstəxana adını görürsə, əlamətlərdən yox, "
        "həmin xəstəxananın tipik xəstələrindən «proqnoz» verə bilər.",
    )

    document.add_heading(
        "21. Böyük dil modellərində tokenləşdirmə, vektor təmsilləri, diqqət, axtarış və hallüsinasiya",
        level=3,
    )
    p(
        document,
        "**LLM** — Large Language Model, yəni böyük dil modelidir (veb səhifədə: iri dil "
        "modeli). **Tokenləşdirmə** mətni modelin emal etdiyi hissələrə bölür; token həmişə "
        "bütöv söz olmur. **Vektor təmsili** (veb səhifədə: yerləşdirmə vektoru, *embedding*) "
        "həmin hissələri ədədlərlə ifadə edir. **Diqqət mexanizmi** kontekstdəki hissələr "
        "arasındakı əlaqələri çəkiləndirir. Xarici mənbədən məlumat gətirilməsi cavab üçün "
        "əlavə əsas verə bilər, lakin düzgünlüyə avtomatik zəmanət yaratmır. Model səlis "
        "davamı yoxlanılmış faktdan üstün tutanda hallüsinasiya qala bilər.",
    )
    example_box(
        document,
        "**Nümunə.** Eyni söz bir və ya bir neçə token ola bilər; buna görə bəzi modellər "
        "qeyri-adi yazılış və ya kodlarda büdrəyir.",
    )

    document.add_heading("Proqram işi", level=3)
    document.add_heading(
        "22. Proqramın yoxlanması, arxitektura, versiyaların idarə edilməsi və davamlı inteqrasiya",
        level=3,
    )
    p(
        document,
        "**Yoxlama** kodun nəzərdə tutulanı edib-etmədiyinə baxır. **Arxitektura** iri "
        "miqyaslı quruluşdur. **Versiya nəzarəti** (çox vaxt Git) kimin nəyi dəyişdiyini "
        "yazır. **Davamlı inteqrasiya** (*Continuous Integration*, CI) komandanın kod "
        "dəyişikliklərini müntəzəm birləşdirib avtomatlaşdırılmış yoxlamalardan keçirməsidir.",
    )
    example_box(
        document,
        "**Nümunə.** Komanda giriş səhifəsini kimin son dəyişdiyini və o dəyişiklikdən sonra "
        "testlərin keçib-keçmədiyini görə bilər.",
    )

    document.add_heading(
        "23. Funksional, məntiqi və eynizamanlı proqramlaşdırma paradiqmaları",
        level=3,
    )
    p(
        document,
        "**Paradiqma** proqramın qurulmasına ümumi yanaşmadır. Funksional üslub funksiya və "
        "dəyərləri, məntiqi üslub münasibət və axtarışı, eynizamanlı və ya paralel üslub isə "
        "vaxtca kəsişən işi önə çıxarır. Eyni məsələ hər üslubda fərqli görünür. "
        "**Eynizamanlılıq** bir neçə işin irəliləməsinin üst-üstə düşməsidir; **paralellik** "
        "işlərin həqiqətən eyni anda icrasıdır. Asinxron yanaşmada nəticə gözlənərkən başqa "
        "iş davam edə bilər.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd; veb səhifədə bu bəndin ayrıca nümunəsi yoxdur).** "
        "Eyni «ən qısa yol» məsələsi funksional dildə funksiya tərkibi, məntiqi dildə "
        "məhdudiyyət, eynizamanlı modeldə isə bir neçə axtarışın üst-üstə düşməsi kimi "
        "görünə bilər.",
    )

    document.add_heading(
        "24. Kompilyatorlar, interpretatorlar və proqramlaşdırma dillərinin semantikası",
        level=3,
    )
    p(
        document,
        "**Interpretator** (veb səhifədə: interpreter) proqramı birbaşa yerinə yetirir. "
        "**Kompilyator** onu əvvəlcə başqa formaya çevirir. **Semantika** dil "
        "konstruksiyalarının nə məna daşıdığını müəyyən edir — proqramın necə yazıldığından "
        "çox, nə etməyə icazəsi olduğu.",
    )
    example_box(
        document,
        "**Nümunə (yenidən işlənmiş sənəd; veb səhifədə bu bəndin ayrıca nümunəsi yoxdur).** "
        "Eyni qısa proqramı əvvəl sətir-sətir izləmək, sonra «bu dil `x = x + 1` deyərkən "
        "dəqiq nə edir?» sualını vermək semantikanı kompilyasiya və ya interpretasiya "
        "mexanizmindən ayırmağa kömək edir.",
    )

    document.add_heading("Hesablama nəzəriyyəsi", level=3)
    document.add_heading("25. Avtomatlar, hesablana bilmə, reduksiya, P, NP və NP-tamlıq", level=3)
    p(
        document,
        "**Avtomat** vəziyyət və keçidlərlə təsvir olunan formal modeldir. "
        "**Hesablana bilmə** məsələnin alqoritmlə həllinin mümkün olub-olmadığını araşdırır. "
        "**Reduksiya** (veb səhifədə: ixtisar) bir məsələni başqa məsələ vasitəsilə həll "
        "etməyə uyğun çevirmədir. **P** (*Polynomial Time*) polinomial vaxtda həll edilən "
        "qərar məsələləridir. **NP** (*Nondeterministic Polynomial Time*) üçün müsbət "
        "cavabın uyğun sübutu polinomial vaxtda yoxlana bilir. NP «polinomial olmayan» "
        "demək deyil. **NP-tam** məsələ NP-dədir və bütün NP məsələləri ona polinomial "
        "vaxtda reduksiya edilə bilir.",
    )
    example_box(
        document,
        "**Nümunə.** A məsələsinin hər nümunəsini B məsələsinin nümunəsinə çevirə bilsən, "
        "B üçün sürətli üsul A üçün də sürətli üsul verər.",
    )


def write_difficult(document: Document) -> None:
    document.add_heading("7. Mövzuların öyrənilməsini çətinləşdirən səbəblər", level=1)
    p(
        document,
        "Mənbələr öyrənənlərin zəif olduğunu iddia etmir. Tədris tədqiqatlarının nəticəsini "
        "göstərir: ideyaların özü asanlıqla səhv oxunur, mütəxəssislər isə öz qısaltmalarını "
        "çox vaxt görmür. Bu nəticələr konkret araşdırmaların auditoriya və metodlarına "
        "bağlıdır. Onlardan yerli auditoriyada yoxlanacaq izah və tapşırıqlar hazırlamaq "
        "üçün istifadə etmək daha məqsədəuyğundur.",
    )

    document.add_heading("7.1 Proqramlaşdırmada üç əsas çətinlik", level=2)
    p(
        document,
        "İlkin proqramlaşdırma üzrə ədəbiyyat icmalı çətinlikləri sintaksis, anlayışların "
        "mənimsənilməsi və məsələ həlli strategiyası ilə əlaqələndirir. **Sintaksis** dilin "
        "yazılış qaydalarıdır. Bu qaydaları bilmək proqramın niyə belə işlədiyini başa "
        "düşmək üçün təkbaşına kifayət etmir. Mənbə Qian və Lehmanın ədəbiyyat icmalına "
        "istinad edir. [6]",
    )
    p(
        document,
        "Əvvəlki araşdırmanı yenidən yoxlayan tədqiqatda müəllimlərin rekursiya, funksiyalar, "
        "dövrlər, səhvlərin tapılması və koddan əvvəl düşünmə və planlaşdırmanı öyrətməkdə "
        "çətinlik çəkdiyi göstərilir. **Təkrar tədqiqat** (*replication*) əvvəlki nəticələrin "
        "başqa vaxtda və ya şəraitdə yenidən müşahidə olunub-olunmadığını yoxlayır. [7]",
    )

    document.add_heading("7.2 Proqramın işləməsi haqqında düzgün təsəvvür", level=2)
    p(
        document,
        "Öyrənən kompüterin əmrləri necə yerinə yetirdiyinə dair düzgün zehni model "
        "qurmadıqda dəyişənlər, dövrlər, funksiyalar, istinadlar və rekursiya bir-biri ilə "
        "əlaqəsi zəif qaydalar kimi görünə bilər. **Notional machine** proqramın icrasını "
        "izah edən sadələşdirilmiş tədris modelidir. İcra ardıcıllığını, yaddaşdakı "
        "qiymətləri və çağırışları görünən etmək bu məqsədə xidmət edir. Mənbə Sorvanın "
        "tədqiqatına və dəyişən qiymətləndirilməsi üzrə işə istinad edir. [8, 9]",
    )

    document.add_heading("7.3 Səhvlərin tapılması və düzəldilməsi", level=2)
    p(
        document,
        "Təcrübəli proqramçı səhv axtararkən çox vaxt özünün artıq avtomatlaşdırdığı "
        "düşüncə addımlarından istifadə edir. Bu addımlar açıq izah edilmədikdə yeni "
        "başlayan yalnız hazır düzəlişi görür. Tədrisdə ehtimalın necə qurulduğunu, hansı "
        "yoxlamanın seçildiyini və nəticənin necə şərh edildiyini göstərmək vacibdir. [10]",
    )

    document.add_heading("7.4 Ayrı sahələrdə müəyyən edilmiş çətinliklər", level=2)
    p(
        document,
        "Aşağıdakı siyahı hər iki mənbədə adı çəkilən sənədləşdirilmiş çətinlik sahələrini "
        "birləşdirir. Hər bənddə qısa izah və ehtiyat saxlanılıb.",
    )
    bullets(
        document,
        [
            "**Verilənlər strukturları və alqoritmlər.** 2026-cı il sistematik icmalında "
            "92 yanlış təsəvvür və öyrənmə çətinliyi müəyyənləşdirilib. Bu say bütün "
            "tələbələrin eyni səhvləri etdiyini deyil, icmalın müxtəlif tədqiqatlarda belə "
            "problemlər topladığını göstərir. [11]",
            "**Obyektyönlü proqramlaşdırma.** Çaşqınlıq çox vaxt sinif ilə obyektin fərqi, "
            "konstruktorlar, obyektlərin bir-birindən qurulması və icra axını ətrafındadır. [12]",
            "**Rekursiya.** İç-içə çağırışların izlənməsi, çağırışların geri qayıtması və "
            "dayanma şərti çətinlik yaradır. Xətti olmayan icranı vizuallaşdırmaq xüsusilə "
            "çətindir. [13]",
            "**İstinadlar və göstəricilər.** Dolayı müraciət, təxəllüs, görünmə sahəsi və "
            "parametr ötürülməsi bir neçə illik təhsildən sonra da çətin qala bilər. [14]",
            "**Eynizamanlılıq və paralellik.** Mümkün icra ardıcıllıqlarının müxtəlifliyi "
            "səhv proqramın bəzən düzgün görünməsinə səbəb olur; yarış və sinxronlaşdırmanı "
            "mühakimə etmək ona görə çətindir. [15]",
            "**SQL.** Genişmiqyaslı təhlildə cədvəllərin birləşdirilməsi, alt sorğular və "
            "GROUP BY ilə bağlı səhvlər qeyd olunur. [16]",
            "**Əməliyyat sistemləri.** Mexanizmi əzbərləmək alternativ həllərin üstünlük və "
            "çatışmazlıqlarını izah etməyə həmişə kifayət etmir. [17]",
            "**Ədəd təqdimatı.** Mövqeli yazılış, ikiyə tamamlayıcı kod və daşma ilə bağlı "
            "yanlış təsəvvürlər kompüterin təşkili kursundan sonra da qala bilər. "
            "**Daşma** nəticənin ayrılmış bit sayına sığmamasıdır. [18]",
            "**Kibertəhlükəsizlik.** Zəiflik, təhdid, risk və qorunma tədbiri qarışdırıla "
            "bilər. Hücum edən tərəfin mümkün addımlarını düşünmək ayrıca bacarıq tələb "
            "edir. [19]",
            "**Dinamik proqramlaşdırma.** Üsulun nə vaxt seçilməsi, altməsələlər arasındakı "
            "əlaqənin qurulması və səmərəsiz təkrarlardan qaçılması çətindir. [20]",
            "**Süni intellekt.** Məktəbli təsəvvürlərini araşdıran tədqiqat uşaqların "
            "sistemlərə insan xüsusiyyətləri aid edə və təlim verilənlərinin rolunu "
            "nəzərdən qaçıra bildiyini göstərir. [21]",
        ],
    )
    p(
        document,
        "Müsabiqə materialını yazanlar bu nəticələri tapşırıq kimi götürə bilər: yalnız "
        "tərifi yox, əskik zehni modeli izah edin.",
    )


def write_first_series(document: Document) -> None:
    document.add_heading("8. İlk nəşrlər üçün təklif olunan plan", level=1)
    document.add_heading("8.1 On iki mövzudan ibarət başlanğıc silsiləsi", level=2)
    p(
        document,
        "Əlçatan tədris məqalələri və izahlı materiallar hazırlamaq üçün aşağıdakı "
        "ardıcıllıq təklif olunur. Başlıqlar ən yüksək prioritet cədvəllə üst-üstə düşür "
        "və qeyri-mütəxəssisin sualı görməsi üçün yazılıb. Bu, nəşr üçün başlanğıc "
        "planıdır; müsabiqənin məcburi mövzu və ya ardıcıllıq tələbi deyil. Auditoriyanın "
        "ehtiyacına görə mövzuların sırası dəyişdirilə və hər mövzu daha kiçik materiallara "
        "bölünə bilər.",
    )
    numbered(
        document,
        [
            "Generativ süni intellekt və ChatGPT tipli sistemlər necə işləyir",
            "Süni intellekt niyə inandırıcı görünən yanlış cavablar verir",
            "Proqram işləyərkən kompüterdə nə baş verir",
            "Dəyişənlər və `x = x + 1` əmrinin mənası",
            "Kod yazmazdan əvvəl alqoritmi necə qurmaq olar",
            "Dövrlər necə işləyir və sonsuz dövr niyə yaranır",
            "Funksiyalar, parametrlər, ekrana çıxarılan nəticələr və qaytarılan qiymətlər",
            "Proqram səhvlərini tapmaq və düzəltmək üçün praktik üsul",
            "İnternet veb səhifəni istifadəçiyə necə çatdırır",
            "Kompüter mətn, təsvir və səsi necə təqdim edir",
            "Parollar, şifrələmə və rəqəmsal imzalar necə işləyir",
            "Verilənlər necə yanlış nəticəyə gətirə bilər və qərəzli təhlili necə tanımaq olar",
        ],
    )

    document.add_heading("8.2 Nümunələr üçün proqramlaşdırma dilinin seçilməsi", level=2)
    p(
        document,
        "Hər iki mənbə eyni üç istiqaməti göstərir: proqramlaşdırma, verilənlər elmi və "
        "süni intellekt nümunələrində Python; veb nümunələrində JavaScript və TypeScript; "
        "verilənlər bazası mövzularında SQL. Aydın izah başqa dildən də istifadə edə bilər, "
        "əgər həmin dil seçilmiş auditoriyanın həqiqətən rastlaşdığı dildirsə.",
    )
    note_box(
        document,
        "**Fərq, qətiyyət.** Yenidən işlənmiş sənəd Python-u «praktik başlanğıc seçimi» "
        "adlandırır və bunu bütün auditoriyalar üçün vahid «ən yaxşı dil» hökmü kimi qəbul "
        "etməməyə çağırır. Veb səhifə isə Python-u «hazırda ən uyğun ümumi tədris dili» "
        "sayır. Bu bələdçi hər iki ifadənı saxlayır və dili mövzunun məqsədinə, "
        "auditoriyanın əvvəlki biliyinə və nümunəni işlətmək imkanına görə seçməyi tövsiyə "
        "edir.",
    )
    p(
        document,
        "2025-ci il Stack Overflow sorğusunda Python istifadəsinin artması və GitHub "
        "hesabatında Python-un süni intellekt layihələrində geniş iştirakı bu dilin "
        "aktuallığını göstərir. TypeScript-in GitHub-dakı yüksəlişi veb nümunələrinə "
        "marağı da əsaslandıra bilər. [4, 5]",
    )


def write_proposal(document: Document) -> None:
    document.add_heading("9. Mövzu təklifinin və izahın hazırlanması", level=1)
    p(
        document,
        "Bu bölmə yalnız yenidən işlənmiş sənəddə var və veb səhifədəki «ilk on iki "
        "məqalə» tövsiyəsini tamamlayır. Hər mövzu üçün aşağıdakı beş suala qısa cavab "
        "hazırlamaq təklif olunur:",
    )
    bullets(
        document,
        [
            "**Auditoriya** — material kimlər üçündür və hansı ilkin bilikləri tələb edir;",
            "**Təhsil əhəmiyyəti** — hansı anlayış və ya bacarığı öyrədir;",
            "**Praktik əlaqə** — həmin bilik harada istifadə olunur;",
            "**Öyrənmə çətinliyi** — hansı yanlış təsəvvürü və ya anlaşılmazlığı aradan qaldırır;",
            "**Gözlənilən nəticə** — öyrənən materialdan sonra nəyi izah və ya tətbiq edə biləcək.",
        ],
    )
    p(
        document,
        "Məsələn, «Funksiya ilə nəticəni ekrana çıxarmaq arasındakı fərq» mövzusu ilkin "
        "proqramlaşdırma öyrənənlərə ünvanlana bilər. Materialın məqsədi qaytarılan "
        "qiymətin başqa hesablamada istifadə oluna bildiyini göstərməkdir. Biri yalnız "
        "ekrana yazan, digəri isə qiymət qaytaran iki qısa funksiya bu fərqi görünən edə bilər.",
    )
    p(
        document,
        "İzahı gündəlik sual və ya kiçik problemlə başlamaq, əsas anlayışları təqdim etmək, "
        "nümunəni addımlarla açmaq və sonda qısa yoxlama tapşırığı vermək olar. Bu quruluş "
        "tövsiyədir və mövzuya uyğun dəyişdirilə bilər.",
    )
    p(
        document,
        "Mövzunun hər üç mənbə qrupunda eyni vaxtda önə çıxması şərt deyil. Rekursiya "
        "gündəlik texnologiya xəbərlərində az görünsə də, onun aydın izahı ciddi təhsil "
        "dəyəri daşıya bilər. Əsas meyar seçilmiş auditoriyanın konkret öyrənmə ehtiyacına "
        "cavab verməkdir.",
    )


def write_related(document: Document) -> None:
    document.add_heading("10. Müsabiqənin əlaqəli səhifələri", level=1)
    p(
        document,
        "Bu bölmə veb səhifədəki «Müsabiqənin digər səhifələri» hissəsindən gəlir. "
        "Yenidən işlənmiş sənəd eyni dəvət səhifəsinə yerli işçi ünvanla istinad edir "
        "(`http://127.0.0.1:8010/az/complex-topics.html`). Həmin ünvan yalnız uyğun yerli "
        "serverdə açılır. Aşağıda ictimai keçidlər verilir.",
    )
    bullets(
        document,
        [
            f"Açıq dəvət: [Çətin mövzu, aydın izah]({SITE_AZ}/complex-topics.html)",
            f"İşçi plan: [Müsabiqənin təşkili]({SITE_AZ}/complex-topics-approach.html)",
            f"Mövzu bələdçisinin veb nüsxəsi: [İnformatika mövzuları]({SITE_AZ}/complex-topics-informatics.html)",
            f"Qısa səsli təqdimat: [Təqdimat]({SITE_AZ}/complex-topics-presentation.html)",
        ],
    )
    p(
        document,
        "Flayer və işçi plan təsdiqlənmiş Əsasnamə deyil. Bu bələdçi də mövzu seçiminə "
        "kömək edən işçi sənəddir.",
    )


def write_glossary(document: Document) -> None:
    document.add_heading("11. Terminlər və ixtisarlar", level=1)
    p(
        document,
        "Bu lüğət mətndə və mövzu siyahılarında işlənən əsas texniki terminləri qısa "
        "şəkildə izah edir. Təşkilat adları və təhsil çərçivələri 3-cü bölmədə açıqlanır. "
        "Veb səhifədə ayrıca lüğət yoxdur; aşağıdakı izahlar yenidən işlənmiş sənəddəndir. "
        "Sinonimlər mötərizədə göstərilib.",
    )
    entries = [
        (
            "AI və SI",
            "Artificial Intelligence və süni intellekt. İnsan zəkası ilə əlaqələndirilən "
            "tapşırıqları yerinə yetirən hesablama üsulları və sistemləri.",
        ),
        (
            "API",
            "Application Programming Interface. Proqramların bir-biri ilə məlumat və əmrlər "
            "mübadiləsi üçün interfeys.",
        ),
        ("CPU", "Central Processing Unit. Əmrləri icra edən mərkəzi prosessor."),
        ("DNS", "Domain Name System. Domen adlarını IP ünvanları ilə əlaqələndirən sistem."),
        ("HTTP", "Hypertext Transfer Protocol. Vebdə sorğu və cavab mübadiləsi üçün protokol."),
        (
            "IP",
            "Internet Protocol. Şəbəkələr arasında paketlərin ünvanlanması və ötürülməsi üçün protokol.",
        ),
        (
            "IoT",
            "Internet of Things. Sensor və qurğuların şəbəkə vasitəsilə məlumat mübadiləsi "
            "apardığı Əşyaların İnterneti.",
        ),
        ("SQL", "Structured Query Language. Verilənlər bazalarında məlumatla işləmək üçün sorğu dili."),
        ("LLM", "Large Language Model. Mətnlə işləyən böyük dil modeli."),
        (
            "OOP",
            "Object-Oriented Programming. Verilənləri və onlarla işləyən əməliyyatları "
            "obyektlər ətrafında quran proqramlaşdırma yanaşması.",
        ),
        ("DSA", "Data Structures and Algorithms. Verilənlər strukturları və alqoritmlər."),
        (
            "CI",
            "Kontekstdən asılı olaraq Continuous Integration — davamlı inteqrasiya və ya "
            "Concept Inventory — anlayışları ölçən diaqnostik test.",
        ),
        (
            "CS1 və SCS1",
            "CS1 ilkin kompüter elmləri və ya proqramlaşdırma kursu üçün işlənən addır. "
            "SCS1 — Second CS1 həmin səviyyədə anlayışları ölçən qiymətləndirmə alətinin adıdır.",
        ),
        (
            "Big-O / Böyük-O",
            "Girişin ölçüsü artdıqca resurs tələbinin artımına asimptotik yuxarı sərhəd verən işarələmə.",
        ),
        (
            "P və NP",
            "P — Polynomial Time; NP — Nondeterministic Polynomial Time. P polinomial "
            "vaxtda həll edilən qərar məsələləridir. NP üçün müsbət cavabın uyğun sübutu "
            "polinomial vaxtda yoxlana bilir. NP «polinomial olmayan» demək deyil.",
        ),
        (
            "NP-tamlıq və reduksiya",
            "NP-tam məsələ NP-dədir və bütün NP məsələləri ona polinomial vaxtda reduksiya "
            "edilə bilir. Reduksiya bir məsələni başqa məsələ vasitəsilə həll etməyə uyğun "
            "çevirmədir. Veb səhifə eyni anlayışı «ixtisar» adlandırır.",
        ),
        (
            "Bul məntiqi",
            "Doğru və yanlış qiymətlərinə, eləcə də VƏ, VƏ YA, DEYİL kimi əməliyyatlara əsaslanan məntiq.",
        ),
        (
            "Görünmə sahəsi və yaşam müddəti",
            "Görünmə sahəsi (əhatə dairəsi) adın harada əlçatan olduğunu, yaşam müddəti "
            "(ömür) isə obyektin və ya saxlanılan məlumatın nə qədər mövcud olduğunu bildirir.",
        ),
        (
            "Parametr və qaytarılan qiymət",
            "Parametr funksiyanın girişini qəbul edir. Qaytarılan qiymət funksiyanın "
            "nəticəsini çağıran hissəyə ötürür.",
        ),
        (
            "İstinad, göstərici və aliasing",
            "İstinad və göstərici obyektə və ya yaddaşa müraciət vasitələridir; dəqiq "
            "davranış dildən asılıdır. Aliasing (təxəllüs) eyni obyektə bir neçə ad və ya "
            "istinadla çıxışdır.",
        ),
        (
            "Stek və heap",
            "Stek son daxil olanın ilk çıxması prinsipi ilə işləyən strukturdur; çağırış "
            "steki funksiyaların icrasını izləyir. Heap yaddaş bölgəsi kimi dinamik ayrılan "
            "yaddaşı, verilənlər strukturu kimi isə ayrıca yığın strukturunu bildirə bilər.",
        ),
        (
            "Rekursiya və baza halı",
            "Rekursiya funksiyanın birbaşa və ya dolayı yolla özünü çağırmasıdır. Baza "
            "halı yeni çağırışa ehtiyac olmayan dayanma vəziyyətidir.",
        ),
        (
            "İrsiyyət və polimorfizm",
            "İrsiyyət siniflər arasında xüsusiyyət və davranışların ötürülməsidir. "
            "Polimorfizm (çoxformalıq) eyni interfeysin müxtəlif obyektlərdə fərqli davranış "
            "verməsinə imkan yaradır.",
        ),
        (
            "Normallaşdırma",
            "Relyasion bazada təkrarlanma və yenilənmə problemlərini azaltmaq üçün cədvəllərin təşkil edilməsi.",
        ),
        (
            "JOIN və aqreqasiya",
            "JOIN əlaqəli cədvəllərin məlumatını birləşdirir. Aqreqasiya çoxlu qiymətdən "
            "cəm, say və orta kimi ümumi nəticə çıxarır.",
        ),
        (
            "Alt sorğu və pəncərə funksiyası",
            "Alt sorğu başqa sorğunun içində yerləşir. Pəncərə funksiyası əlaqəli sətirlər "
            "üzrə hesab apararaq ayrı sətirləri nəticədə saxlayır.",
        ),
        (
            "Tranzaksiya və izolyasiya",
            "Tranzaksiya əlaqəli əməliyyatları vahid iş kimi idarə edir. İzolyasiya "
            "(izolə) paralel tranzaksiyaların bir-birinin aralıq nəticələrini necə görə "
            "bildiyini müəyyənləşdirir.",
        ),
        (
            "Eynizamanlılıq və paralellik",
            "Eynizamanlılıq bir neçə işin irəliləməsinin üst-üstə düşməsidir. Paralellik "
            "işlərin həqiqətən eyni anda icrasıdır. Asinxron yanaşmada nəticə gözlənərkən "
            "başqa iş davam edə bilər.",
        ),
        (
            "Yarış vəziyyəti və kilidlənmə",
            "Yarış vəziyyətində nəticə icra ardıcıllığından asılı olur. Qarşılıqlı "
            "kilidlənmə (tıxac) işlərin bir-birinin saxladığı resursları gözləyərək "
            "dayanmasıdır.",
        ),
        (
            "Proses, axın və virtual yaddaş",
            "Proses icra olunan proqramın mühitidir; axın həmin mühitdə icra yoludur. "
            "Virtual yaddaş proqramın gördüyü ünvanları fiziki yaddaşdan ayırır.",
        ),
        (
            "Autentifikasiya və avtorizasiya",
            "Autentifikasiya kimliyi, avtorizasiya isə əməliyyatı etməyə icazəni yoxlayır.",
        ),
        (
            "Fişinq və iki amilli autentifikasiya",
            "Fişinq aldadıcı mesaj və ya saytla məlumat əldə etməyə cəhddir. İki amilli "
            "autentifikasiya (iki faktorlu təsdiq) fərqli kateqoriyalardan iki yoxlama "
            "vasitəsindən istifadə edir.",
        ),
        (
            "Şifrələmə, heşləmə və imza",
            "Şifrələmə məlumatı açarla qoruyur. Heşləmə qısa rəqəmsal iz yaradır. "
            "Rəqəmsal imza məlumatın bütövlüyünü və imzalayanla əlaqəsini yoxlamağa imkan verir.",
        ),
        (
            "Təlim və qiymətləndirmə",
            "Təlim zamanı model nümunələrdən öyrənir. Qiymətləndirmə onun ayrıca məlumat "
            "üzərində necə işlədiyini yoxlayır.",
        ),
        (
            "Həddindən artıq və yetərsiz öyrənmə",
            "Overfitting zamanı model təlim məlumatının özəlliklərinə həddən çox uyğunlaşır. "
            "Underfitting zamanı model əsas əlaqələri kifayət qədər öyrənmir.",
        ),
        (
            "Sinif balanssızlığı və qərəz",
            "Sinif balanssızlığı nümunə kateqoriyalarının sayca qeyri-bərabər olmasıdır. "
            "Qərəz nəticələrin müəyyən qruplar və ya şərait üçün sistemli şəkildə təhrif "
            "olunması ola bilər.",
        ),
        (
            "İzah edilə bilmə",
            "Modelin nəticəsinə təsir edən amillərin insanlar üçün başadüşülən şəkildə açıqlana bilməsi.",
        ),
        (
            "Token, vektor təmsili və diqqət",
            "Token mətnin emal vahididir; həmişə bütöv söz olmur. Vektor təmsili onu "
            "ədədlərlə göstərir. Diqqət mexanizmi kontekst hissələri arasındakı əlaqələri "
            "çəkiləndirir.",
        ),
        (
            "Git və GitHub",
            "Git dəyişiklik tarixçəsini izləyən versiya idarəetmə sistemidir. GitHub Git "
            "layihələrinin saxlanması və birgə işlənməsi üçün platformadır.",
        ),
        (
            "İkiyə tamamlayıcı kod",
            "Sabit bit sayında işarəli tam ədədlərin, o cümlədən mənfi ədədlərin təqdim edilməsi üsulu.",
        ),
        (
            "Avtomat və hesablana bilmə",
            "Avtomat vəziyyət və keçidlərlə təsvir olunan formal modeldir. Hesablana bilmə "
            "məsələnin alqoritmlə həllinin mümkün olub-olmadığını araşdırır.",
        ),
        (
            "Müştəri və server",
            "Müştəri xidmət üçün sorğu göndərir, server həmin sorğunu emal edib cavab verir.",
        ),
        (
            "İnsan xüsusiyyətləri aid etmə",
            "Sistemə insan kimi düşüncə, niyyət və hisslər aid etmək. Süni intellekt "
            "barədə izahlarda bu fərziyyə mexanizmin başa düşülməsini çətinləşdirə bilər.",
        ),
    ]
    rows = [["Termin", "İzah"]]
    rows.extend([[name, text] for name, text in entries])
    add_table(document, rows)


def write_references(document: Document) -> None:
    document.add_heading("12. Mənbələr və istinadlar", level=1)
    p(
        document,
        "Mətn daxilindəki nömrələr aşağıdakı mənbələrə uyğundur. Eyni mənbəyə verilmiş "
        "təkrar keçidlər birləşdirilib. Tədqiqatlar üçün mənbə sənədindəki qısa təsviri "
        "adlar saxlanılıb; onlar tam biblioqrafik başlıq kimi təqdim edilmir. "
        "Keçidlər kliklənəndir.",
    )

    refs = [
        (
            "1",
            "CSTA’s 2026 PK–12 Standards",
            "https://csteachers.org/pk12standards/",
        ),
        (
            "2",
            "ACM/IEEE-CS/AAAI CS2023 — bilik sahələri",
            "https://csed.acm.org/knowledge-areas/",
        ),
        (
            "3",
            "AP Computer Science Principles",
            "https://apcentral.collegeboard.org/courses/ap-computer-science-principles",
        ),
        (
            "4",
            "2025 Stack Overflow Developer Survey",
            "https://survey.stackoverflow.co/2025/technology",
        ),
        (
            "5",
            "GitHub Octoverse 2025",
            "https://github.blog/news-insights/octoverse/octoverse-a-new-developer-joins-github-every-second-as-ai-leads-typescript-to-1/",
        ),
        ("6", "Qian and Lehman’s literature review", "https://doi.org/10.1145/3077618"),
        ("7", "20-year CS1 replication study", "https://doi.org/10.1145/3730405"),
        ("8", "Sorva’s notional-machine research", "https://doi.org/10.1145/2483710.2483713"),
        ("9", "Study of variable evaluation", "https://doi.org/10.1145/3017680.3017724"),
        ("10", "Systematic review of debugging instruction", "https://doi.org/10.1145/3690652"),
        ("11", "DSA systematic review", "https://doi.org/10.1080/08993408.2026.2633989"),
        ("12", "Long-term OOP study", "https://doi.org/10.1080/08993400500224310"),
        ("13", "Recursion concept inventory", "https://doi.org/10.1080/08993408.2017.1414728"),
        ("14", "Indirection study", "https://doi.org/10.15388/infedu.2507.015"),
        ("15", "Concurrency study", "https://doi.org/10.15388/infedu.2021.29"),
        ("16", "SQL misconception study", "https://doi.org/10.1145/2899415.2899464"),
        ("17", "Operating-systems concept inventory", "https://doi.org/10.1145/2538862.2538886"),
        ("18", "Number-representation study", "https://doi.org/10.1080/08993408.2011.611712"),
        (
            "19",
            "Cybersecurity misconception study",
            "https://digitalcommons.kennesaw.edu/jcerp/vol2018/iss1/5/",
        ),
        (
            "20",
            "Replication study on dynamic programming",
            "https://doi.org/10.1080/08993408.2022.2079865",
        ),
        ("21", "Study of pupils’ AI conceptions", "https://doi.org/10.1016/j.caeai.2022.100095"),
        (
            "22",
            "Müsabiqənin açıq dəvəti (ictimai ünvan). Yenidən işlənmiş sənəddə yerli "
            "inkişaf ünvanı da göstərilib: http://127.0.0.1:8010/az/complex-topics.html",
            f"{SITE_AZ}/complex-topics.html",
        ),
        ("23", "CS2023 rəsmi səhifəsi", "https://csed.acm.org/"),
        (
            "24",
            "CS2023 bilik və səriştə modeli",
            "https://csed.acm.org/wp-content/uploads/2024/04/1.3-Introduction-to-Knowledge-Model.pdf",
        ),
        ("25", "Sorğunun metodologiyası", "https://survey.stackoverflow.co/2025/methodology"),
        ("26", "Sorğunun süni intellekt bölməsi", "https://survey.stackoverflow.co/2025/ai"),
        (
            "27",
            "Dinamik proqramlaşdırma üzrə konsepsiya inventarı tədqiqatı",
            "https://arxiv.org/abs/2411.14655",
        ),
        (
            "28",
            "SCS1-in yoxlanılması və tətbiqi",
            "https://www.sci.sdsu.edu/crmse/msed/papers/parker2-parker-guzdial-engleman_SCS1.pdf",
        ),
        (
            "29",
            "İnformatika mövzularının Azərbaycanca veb bələdçisi",
            f"{SITE_AZ}/complex-topics-informatics.html",
        ),
        (
            "30",
            "Müsabiqənin təşkili — işçi plan",
            f"{SITE_AZ}/complex-topics-approach.html",
        ),
        (
            "31",
            "Müsabiqənin qısa səsli təqdimatı",
            f"{SITE_AZ}/complex-topics-presentation.html",
        ),
    ]
    for num, title, url in refs:
        para = document.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        add_formatted_run(para, f"[{num}] {title} — ")
        add_hyperlink(para, url, url)

    p(
        document,
        "Bu inteqrasiya olunmuş nəşrin özü yeni mənbə deyil. O, yuxarıdakı sənədləri və "
        "iki daxili mənbəni — yenidən işlənmiş Word sənədini və Azərbaycanca veb səhifəni — "
        "bir araya gətirir.",
    )


def build() -> Path:
    sources = [
        ROOT / "documents" / "Popular_Topics_in_Informatics_Reworked_AZ.docx",
        ROOT / "az" / "complex-topics-informatics.html",
    ]
    missing = [str(s) for s in sources if not s.is_file()]
    if missing:
        raise FileNotFoundError("Source missing: " + "; ".join(missing))

    document = Document()
    setup_styles(document)
    setup_az_header_footer(
        document,
        "İnformatika üzrə aktual mövzular — inteqrasiya olunmuş bələdçi",
    )
    set_update_fields_on_open(document)

    add_title_page(document)
    add_toc(document)
    write_intro(document)
    write_discrepancies(document)
    write_sources(document)
    write_priority(document)
    write_school(document)
    write_university(document)
    write_difficult(document)
    write_first_series(document)
    write_proposal(document)
    write_related(document)
    write_glossary(document)
    write_references(document)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes)")
