# -*- coding: utf-8 -*-
"""
Ensure every EN Word CV has an Awards, honors and fellowships table.

- Creates the section + 3-column table when missing
- Populates curated rows drawn from the CV / web profile (no invented facts)
- Leaves a header-only table when no awards are known

Usage:
  python helpers/_cv_ensure_awards_tables.py --dry-run
  python helpers/_cv_ensure_awards_tables.py
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _cv_en_polish_lib import (
    EN_CV_DIR,
    PUBLICATION_HEADINGS,
    extract_cv,
    iter_blocks,
)
from _normalize_cv_contact_links import set_cell_text

SUMMARY = Path(__file__).with_name("_cv_ensure_awards_summary.json")

# Curated awards for CVs that currently lack an Awards table.
# Only facts already present in the Word CV and/or EN web profile.
AWARDS_FOR_MISSING: dict[str, list[dict[str, str]]] = {
    "Applebaum_Mark_CV_EN.docx": [
        {
            "year": "2022",
            "honor": "Elected to the International Committee of the International Group for Mathematical Creativity and Giftedness (IGMCG) and to its Education Committee",
            "body": "IGMCG",
        },
    ],
    "Allahverdiyev_Oruc_CV_EN.docx": [
        {
            "year": "2014",
            "honor": "Honor Student distinction on completion of doctoral studies",
            "body": "",
        },
    ],
    "Jannataliyev_Rasim_CV_EN.docx": [
        {
            "year": "2008–2009",
            "honor": "DAAD-funded research stay (ecological problems of the Caspian Sea)",
            "body": "DAAD; University of Rostock, Germany",
        },
    ],
    "Mammadov_Agamali_CV_EN.docx": [
        {
            "year": "",
            "honor": "Russian President’s Culture Award",
            "body": "Russian Federation",
        },
        {
            "year": "",
            "honor": "Education-Online Grant award",
            "body": "V. Potanin Foundation",
        },
    ],
    "Osmanli_Zarifa_CV_EN.docx": [
        {
            "year": "February–March 2023",
            "honor": "Marie Skłodowska-Curie RISE secondment (REFRACT)",
            "body": "EU Marie Skłodowska-Curie RISE; National University of Quilmes, Argentina",
        },
        {
            "year": "2014–2017",
            "honor": "Three-year Turkey scholarship (MSc studies)",
            "body": "Turkey",
        },
        {
            "year": "2009–2013",
            "honor": "Four-year scholarship",
            "body": "Ministry of Education of Azerbaijan",
        },
    ],
    "Salamov_Asaf_CV_EN.docx": [
        {
            "year": "2014–2020",
            "honor": "Listed among the world’s most cited scientists",
            "body": "Clarivate Analytics",
        },
    ],
}


def save_doc(path: Path, doc: Document) -> None:
    tmp = path.with_name(f"__tmp_{path.name}")
    try:
        doc.save(str(tmp))
        tmp.replace(path)
    except OSError:
        alt = path.with_name(f"__tmp2_{path.name}")
        doc.save(str(alt))
        alt.replace(path)


def find_awards_table(doc: Document) -> Table | None:
    blocks = list(iter_blocks(doc))
    for i, (kind, obj) in enumerate(blocks):
        if kind != "p":
            continue
        text = (obj.text or "").strip()
        if text.lower().startswith("awards"):
            if i + 1 < len(blocks) and blocks[i + 1][0] == "t":
                return blocks[i + 1][1]
    return None


def ensure_awards_table(doc: Document) -> Table:
    existing = find_awards_table(doc)
    if existing is not None:
        return existing

    blocks = list(iter_blocks(doc))
    insert_before = None
    for kind, obj in blocks:
        if kind != "p":
            continue
        text = (obj.text or "").strip()
        if PUBLICATION_HEADINGS.match(text) or text.lower().startswith(
            "date of last update"
        ):
            insert_before = obj._p
            break

    body = doc.element.body
    h = doc.add_paragraph("Awards, honors and fellowships")
    h_el = h._p
    body.remove(h_el)
    if insert_before is not None:
        insert_before.addprevious(h_el)
    else:
        body.append(h_el)

    table = doc.add_table(rows=1, cols=3)
    tbl_el = table._tbl
    body.remove(tbl_el)
    h_el.addnext(tbl_el)
    set_cell_text(table.rows[0].cells[0], "Year")
    set_cell_text(table.rows[0].cells[1], "Award / honor")
    set_cell_text(table.rows[0].cells[2], "Awarding body")
    return table


def clone_row(table: Table, row_idx: int = 0):
    row = table.rows[row_idx]
    new_tr = deepcopy(row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def award_exists(table: Table, year: str, honor: str) -> bool:
    tokens = [t for t in re.findall(r"[a-z0-9]{4,}", honor.lower())][:6]
    for row in table.rows[1:]:
        row_text = " ".join(c.text.strip().lower() for c in row.cells)
        if year and year[:4] in row_text and tokens:
            if sum(1 for t in tokens if t in row_text) >= min(2, len(tokens)):
                return True
        if tokens and sum(1 for t in tokens if t in row_text) >= min(3, len(tokens)):
            return True
    return False


def sort_awards(table: Table) -> None:
    if len(table.rows) <= 2:
        return
    rows = list(table.rows)[1:]

    def key(r):
        y = (r.cells[0].text or "").strip()
        m = re.search(r"((?:19|20)\d{2})", y)
        return int(m.group(1)) if m else 0

    ordered = sorted(rows, key=key, reverse=True)
    for r in rows:
        table._tbl.remove(r._tr)
    for r in ordered:
        table._tbl.append(r._tr)


def add_awards(table: Table, awards: list[dict[str, str]]) -> list[str]:
    changes = []
    for a in awards:
        year = a.get("year") or ""
        honor = a.get("honor") or ""
        body = a.get("body") or ""
        if not honor:
            continue
        if award_exists(table, year, honor):
            continue
        row = clone_row(table, 0)
        set_cell_text(row.cells[0], year)
        set_cell_text(row.cells[1], honor)
        set_cell_text(row.cells[2], body)
        changes.append(f"{year}|{honor[:70]}")
    if changes:
        sort_awards(table)
    return changes


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    summary = []
    for path in sorted(EN_CV_DIR.glob("*_CV_EN.docx")):
        data = extract_cv(path)
        had = any(
            (t.get("section") or "").lower().startswith("awards") for t in data["tables"]
        )
        if had:
            continue

        awards = AWARDS_FOR_MISSING.get(path.name, [])
        entry = {
            "file": path.name,
            "created_table": True,
            "rows_added": [f"{a.get('year','')}|{a.get('honor','')[:70]}" for a in awards],
        }
        if args.dry_run:
            summary.append(entry)
            print("DRY", path.name, f"rows={len(awards)}")
            continue

        doc = Document(str(path))
        table = ensure_awards_table(doc)
        added = add_awards(table, awards) if awards else []
        save_doc(path, doc)
        entry["rows_added"] = added
        summary.append(entry)
        print("OK", path.name, f"rows={len(added)}")

    # verify
    still_missing = []
    for path in sorted(EN_CV_DIR.glob("*_CV_EN.docx")):
        data = extract_cv(path)
        if not any(
            (t.get("section") or "").lower().startswith("awards") for t in data["tables"]
        ):
            # dry-run won't create; only flag after real run
            if not args.dry_run:
                still_missing.append(path.name)

    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"processed={len(summary)} still_missing={len(still_missing)} dry_run={args.dry_run}")
    if still_missing:
        print("MISSING", still_missing)


if __name__ == "__main__":
    main()
