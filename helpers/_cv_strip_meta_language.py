# -*- coding: utf-8 -*-
"""Remove self-referential / meta CV phrasing from EN and AZ Word CVs."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
WORD = {
    "en": ROOT / "en" / "scientists" / "CV" / "Word",
    "az": ROOT / "az" / "scientists" / "CV" / "Word",
}


def set_run_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.text = text


def tidy(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s+([,.;:])", r"\1", text)
    text = re.sub(r"\.\s*\.", ". ", text)
    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"  +", " ", text)
    text = re.sub(r"^\s*[,;]\s*", "", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    # Fix spaces before closing paren leftovers
    text = re.sub(r"\(\s+", "(", text)
    text = re.sub(r"\s+\)", ")", text)
    text = re.sub(r",\s*,", ",", text)
    text = re.sub(r";\s*;", ";", text)
    return text


def clean_en(text: str) -> str:
    original = text

    # Drop sentences that only comment on missing lists / non-printing
    text = re.sub(
        r"(?i)\s*The CV names a list of publications but does not print the references\.?",
        "",
        text,
    )
    text = re.sub(
        r"(?i)\s*without listing them\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i),\s*without listing them",
        "",
        text,
    )
    text = re.sub(
        r"(?i),\s*without naming them\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i),\s*without titles\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i),\s*without proficiency levels\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i),\s*without giving the institutions or thesis titles\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i),\s*without separate (start )?years\.?",
        ".",
        text,
    )
    text = re.sub(
        r"(?i)\s*Master’s and postgraduate study at the same academy are mentioned without separate years\.?",
        "",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV does not give the years( of those degrees)?\.?",
        "",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV does not name a doctorate\.?",
        "",
        text,
    )

    # Equivalence / spelling / title print notes
    text = re.sub(
        r"(?i)\s*The CV equates this with the German Dr\. habil\.?",
        " Equivalent to the German Dr. habil.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV equates this with the German Dr\. and the U\.S\. PhD\.?",
        " Equivalent to the German Dr. and the U.S. PhD.",
        text,
    )
    text = re.sub(
        r"(?i),\s*which the CV equates with the German Dr\. and the U\.S\. PhD",
        ", equivalent to the German Dr. and the U.S. PhD",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV equates that degree with the German Dr\. habil\.?",
        " Equivalent to the German Dr. habil.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV calls this the U\.S\. equivalent of a Master of Arts in physics\.?",
        " Equivalent to a Master of Arts in physics.",
        text,
    )
    text = re.sub(
        r"(?i),\s*described in the CV as equivalent to an MSc,",
        ", equivalent to an MSc,",
        text,
    )
    text = re.sub(
        r"(?i),\s*described as equivalent to a PhD,",
        ", equivalent to a PhD,",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV states that this is equivalent to a PhD\.?",
        " Equivalent to a PhD.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV states that the diploma is equivalent to an MSc\.?",
        " Equivalent to an MSc.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV also calls this degree a PhD\.?",
        " Also awarded as a PhD.",
        text,
    )
    text = re.sub(
        r"(?i),\s*which the CV also calls a PhD,",
        " (also awarded as a PhD),",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV titles this a PhD thesis\.?",
        " Also titled a PhD thesis.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV prints Associated Professor\.?",
        "Associate Professor",
        text,
    )
    text = re.sub(
        r"(?i),\s*a rank the CV spells Associated Professor and dates from 2018",
        " (Associate Professor since 2018)",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV prints B\.S\.?",
        "BSc",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV prints the range 2021–2021\.?",
        "2021",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV title is Prof\. Dr\.\s*",
        "Professor. ",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV heading gives the rank Doç\. Dr\., and the appointment line from 4 January 2012 describes him as yardımcı doçent\.",
        " Rank: Doç. Dr.; appointed yardımcı doçent from 4 January 2012.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV heading gives the degree PhD\.?",
        " Holds a PhD.",
        text,
    )
    text = re.sub(
        r"(?i)\s*The CV gives both labels for 1981\.?",
        " Both master’s and BSc labels are recorded for 1981.",
        text,
    )
    text = re.sub(
        r"(?i)\(program title in the CV:\s*([^)]+)\)",
        r"(\1)",
        text,
    )
    text = re.sub(
        r"(?i),\s*as given in the English text of the CV,",
        ",",
        text,
    )
    text = re.sub(
        r"(?i);\s*the thesis named in the CV is",
        "; thesis:",
        text,
    )
    text = re.sub(
        r"(?i)Thesis topics named in the CV:",
        "Thesis topics:",
        text,
    )
    text = re.sub(
        r"(?i)Named completed theses are in the CV, including",
        "Completed theses include",
        text,
    )
    text = re.sub(
        r"(?i)Completed theses at Hasan Kalyoncu University are listed in the CV",
        "Supervised completed theses at Hasan Kalyoncu University",
        text,
    )
    text = re.sub(
        r"(?i)Completed master’s theses at Kafkas University are listed in the CV",
        "Supervised completed master’s theses at Kafkas University",
        text,
    )
    text = re.sub(
        r"(?i)Named theses at Kafkas University are listed in the CV",
        "Supervised named theses at Kafkas University",
        text,
    )
    text = re.sub(
        r"(?i)The CV lists courses taught in 1989–2024",
        "Taught courses in 1989–2024",
        text,
    )
    text = re.sub(
        r"(?i)\s*The employment record states that the dissertation was submitted in 1990\.?",
        " Dissertation submitted in 1990.",
        text,
    )
    text = re.sub(
        r"(?i),\s*dated 2019–present in the CV\.?",
        " (2019–present).",
        text,
    )

    # Generic "The CV <verb>..." → active / direct
    replacements = [
        (
            r"(?i)Aladdin Allahverdiyev is described as a professor and doctor\.\s*The CV lists him as",
            "Aladdin Allahverdiyev is a professor and doctor. Positions include",
        ),
        (
            r"(?i)The CV describes six patents or inventor’s certificates,",
            "Holds six patents or inventor’s certificates,",
        ),
        (
            r"(?i)The CV also lists earlier study and teaching-assistant work in Strasbourg and Baku\.",
            "Earlier study and teaching-assistant work took place in Strasbourg and Baku.",
        ),
        (
            r"(?i)The CV states an h-index of 4\.",
            "h-index: 4.",
        ),
        (
            r"(?i)The CV dates a PhD defense to 11 April 1998 and a habilitation defense to 3 June 2009\.",
            "PhD defense: 11 April 1998; habilitation defense: 3 June 2009.",
        ),
        (
            r"(?i)The CV lists state assignments and grants from",
            "Holds state assignments and grants from",
        ),
        (
            r"(?i)The CV states more than 150 publications, including 12 textbooks and monographs, and supervision of 10 candidates of sciences and 2 doctors of sciences\.",
            "Published more than 150 works, including 12 textbooks and monographs; supervised 10 candidates of sciences and 2 doctors of sciences.",
        ),
        (
            r"(?i)The CV states more than 450 works, inventions, monographs, and textbooks, supervision of 7 candidates of sciences, and more than 60 talks, including 12 invited plenary lectures\.",
            "Published more than 450 works, inventions, monographs, and textbooks; supervised 7 candidates of sciences; delivered more than 60 talks, including 12 invited plenary lectures.",
        ),
        (
            r"(?i)The CV, dated June 2024, states more than 450 articles, 41 book chapters, 7 patents, and 13 monographs, with more than 30,000 citations and an h-index of 85\.",
            "As of June 2024: more than 450 articles, 41 book chapters, 7 patents, and 13 monographs; more than 30,000 citations; h-index 85.",
        ),
        (
            r"(?i)The CV gives her specialty as physician and ophthalmologist-surgeon\.",
            "Specialty: physician and ophthalmologist-surgeon.",
        ),
        (
            r"(?i)The CV states that she is the author of 18 invention patents, 15 rationalization proposals, and more than 140 scientific articles and abstracts\.",
            "Authored 18 invention patents, 15 rationalization proposals, and more than 140 scientific articles and abstracts.",
        ),
        (
            r"(?i)It also names gold medals of Hippocrates and V\. I\. Vernadsky and the title Honored Worker of Science and Education\.",
            "Awards include gold medals of Hippocrates and V. I. Vernadsky and the title Honored Worker of Science and Education.",
        ),
        (
            r"(?i)The CV lists 131 journal papers and 62 conference papers, with 2,825 citations and an h-index of 28\.",
            "Published 131 journal papers and 62 conference papers; 2,825 citations; h-index 28.",
        ),
        (
            r"(?i)The thesis, as given in the English text of the CV, is",
            "Thesis:",
        ),
        (
            r"(?i)The CV states more than 20 articles and book chapters in maxillofacial surgery and implantology, and two surgical techniques and a suture method that carry his surname\.",
            "Published more than 20 articles and book chapters in maxillofacial surgery and implantology; developed two surgical techniques and a suture method that carry his surname.",
        ),
        (
            r"(?i)The CV lists a master’s degree and a BSc in applied mathematics from Azerbaijan State University in 1981\.",
            "Holds a master’s degree and a BSc in applied mathematics from Azerbaijan State University (1981).",
        ),
        (
            r"(?i)The CV lists 8 books and states that he is an author or coauthor of 186 journal and conference papers\.",
            "Authored or coauthored 8 books and 186 journal and conference papers.",
        ),
        (
            r"(?i)The CV reports 7,192 Google Scholar citations and an h-index of 46\.",
            "Google Scholar: 7,192 citations; h-index 46.",
        ),
        (
            r"(?i)The CV gives a 1981 Promotion A \(Dr\. rer\. nat\.\) on contact problems of plastic flow in a thin metal layer, and a 2000 Promotion B, described as a habilitation and Doctor of Science, on nonstationary contact problems of flow in a thin plastic layer\.",
            "Received a 1981 Promotion A (Dr. rer. nat.) on contact problems of plastic flow in a thin metal layer, and a 2000 Promotion B (habilitation / Doctor of Science) on nonstationary contact problems of flow in a thin plastic layer.",
        ),
        (
            r"(?i)The CV states more than 200 research publications and more than 45 textbooks, and lists 16 key works\.",
            "Published more than 200 research works and more than 45 textbooks, including 16 key works.",
        ),
        (
            r"(?i)The CV lists books of essays, poetry, and Turkish-Estonian language guides, and translations from Azerbaijani, Estonian, and Turkish\.",
            "Published books of essays, poetry, and Turkish-Estonian language guides, and translations from Azerbaijani, Estonian, and Turkish.",
        ),
        (
            r"(?i)The CV states more than 300 scientific works, monographs, and patents\.",
            "Published more than 300 scientific works, monographs, and patents.",
        ),
        (
            r"(?i);\s*the CV says this was recognized in Azerbaijan as Candidate of Legal Sciences\.",
            "; recognized in Azerbaijan as Candidate of Legal Sciences.",
        ),
        (
            r"(?i)\s*The CV says the degree was recognized in Azerbaijan as Candidate of Legal Sciences\.",
            " Recognized in Azerbaijan as Candidate of Legal Sciences.",
        ),
        (
            r"(?i)The CV states 8 monographs, 3 textbooks, 2 teaching aids, more than 250 articles, 3 Scopus articles, and supervision of 25 candidates and 2 doctors of legal sciences\.",
            "Published 8 monographs, 3 textbooks, 2 teaching aids, more than 250 articles, and 3 Scopus articles; supervised 25 candidates and 2 doctors of legal sciences.",
        ),
        (
            r"(?i)The CV lists 8 published articles\.",
            "Published 8 articles.",
        ),
        (
            r"(?i)The CV names one journal article, in the International Journal of Cancer \(2020\), with a DOI\.",
            "Published a journal article in the International Journal of Cancer (2020).",
        ),
        (
            r"(?i)The CV lists selected books, articles, and talks on Iranian Turkic press, language, and identity\.",
            "Published selected books, articles, and talks on Iranian Turkic press, language, and identity.",
        ),
        (
            r"(?i);\s*the CV also gives his title as Dr\.-Ing\.",
            "; title: Dr.-Ing.",
        ),
        (
            r"(?i)The CV lists his languages as German, Azerbaijani, Turkish, and Russian\.",
            "Languages: German, Azerbaijani, Turkish, and Russian.",
        ),
        (
            r"(?i)The CV states a bachelor’s degree in electronics engineering in 1980 and further degrees in Japan: a master’s in 1984 at Toyohashi University of Technology, a doctorate in 1987 at Nagoya University, and a doctorate in 2003 at Tohoku University\.",
            "Holds a bachelor’s degree in electronics engineering (1980) and further degrees in Japan: master’s (1984, Toyohashi University of Technology), doctorate (1987, Nagoya University), and doctorate (2003, Tohoku University).",
        ),
        (
            r"(?i)It also states a first PhD from Toyohashi University of Technology and a second from Tohoku University\.",
            "Holds a first PhD from Toyohashi University of Technology and a second from Tohoku University.",
        ),
        (
            r"(?i)The CV states more than 96 articles, nearly 100 patents, three books in Japanese, and more than 15 book chapters\.",
            "Published more than 96 articles, nearly 100 patents, three books in Japanese, and more than 15 book chapters.",
        ),
        (
            r"(?i)The CV lists books, chapters, and symposium papers, mostly on Azerbaijani history, literature, and education\.",
            "Published books, chapters, and symposium papers, mostly on Azerbaijani history, literature, and education.",
        ),
        (
            r"(?i);\s*the CV says the thesis defense was transferred to Baku State University\.",
            "; thesis defense transferred to Baku State University.",
        ),
        (
            r"(?i)The CV lists a 2016 book and several papers on Elmira Aliyeva and piano education in Azerbaijan\.",
            "Published a 2016 book and several papers on Elmira Aliyeva and piano education in Azerbaijan.",
        ),
        (
            r"(?i)The CV states an h-index of 44, more than 52 monographs and textbooks, and supervision of 25 candidates of sciences and 3 doctors of sciences\.",
            "h-index 44; published more than 52 monographs and textbooks; supervised 25 candidates of sciences and 3 doctors of sciences.",
        ),
        (
            r"(?i)The CV gives 400 publications, an h-index of 33, and 41 candidates and doctors of sciences prepared\.",
            "Published 400 works; h-index 33; prepared 41 candidates and doctors of sciences.",
        ),
        (
            r"(?i)The CV lists four main works, including a 2015 book on Turkey and contributions connected with volumes edited by A\. Bulatov\.",
            "Main works include a 2015 book on Turkey and contributions to volumes edited by A. Bulatov.",
        ),
        (
            r"(?i)The CV states more than 100 articles and about 20 books\.",
            "Published more than 100 articles and about 20 books.",
        ),
        (
            r"(?i)The CV states more than 350 publications\.",
            "Published more than 350 works.",
        ),
        (
            r"(?i)The CV describes a Candidate of Technical Sciences degree from postgraduate study in 1988–1991, and labels that thesis a PhD thesis, and a Doctor of Technical Sciences degree from doctoral study in 1999–2003\.",
            "Received a Candidate of Technical Sciences degree (postgraduate study 1988–1991; also titled a PhD thesis) and a Doctor of Technical Sciences degree (doctoral study 1999–2003).",
        ),
        (
            r"(?i)The CV states more than 390 scientific works, including 18 monographs and 13 textbooks, supervision of 15 candidate and 6 doctoral dissertations, and 4 inventions\.",
            "Published more than 390 scientific works, including 18 monographs and 13 textbooks; supervised 15 candidate and 6 doctoral dissertations; holds 4 inventions.",
        ),
        (
            r"(?i)The CV describes him as a doctor of technical sciences, professor, and Honored Science Worker of the Republic of Azerbaijan\.",
            "Doctor of technical sciences, professor, and Honored Science Worker of the Republic of Azerbaijan.",
        ),
        (
            r"(?i)It states scientific training in the United Kingdom in 1977–1978 and that he rose from assistant to head of department and defended candidate and doctoral dissertations\.",
            "Completed scientific training in the United Kingdom in 1977–1978; rose from assistant to head of department and defended candidate and doctoral dissertations.",
        ),
        (
            r"(?i)The CV states more than 300 scientific works, monographs, textbooks, and articles\.",
            "Published more than 300 scientific works, monographs, textbooks, and articles.",
        ),
        (
            r"(?i)The CV also lists associate-professor and professor ranks\.",
            "Holds associate-professor and professor ranks.",
        ),
        (
            r"(?i)The CV gives his title as Prof\. Dr\.\s*",
            "Title: Prof. Dr. ",
        ),
        (
            r"(?i)It states more than 150 scientific works, including 5 monographs, of which 12 articles are in Web of Science or Scopus journals, and more than 100 conference papers\.",
            "Published more than 150 scientific works, including 5 monographs (12 articles in Web of Science or Scopus) and more than 100 conference papers.",
        ),
        (
            r"(?i)Google Scholar figures in the CV are an h-index of 10 and an i10-index of 11\.",
            "Google Scholar: h-index 10; i10-index 11.",
        ),
        (
            r"(?i)The CV states more than 800 scientific publications, of which 632 are in Science Citation Index journals\.",
            "Published more than 800 scientific works, of which 632 are in Science Citation Index journals.",
        ),
        (
            r"(?i)The CV states more than 40 books, works, and textbooks\.",
            "Published more than 40 books, works, and textbooks.",
        ),
    ]
    for pat, repl in replacements:
        text = re.sub(pat, repl, text)

    # Residual generic cleanups
    text = re.sub(r"(?i)\bthe candidate\b(?!\s+of\b)", "the scholar", text)
    text = re.sub(r"(?i)\bin this CV\b", "", text)
    text = re.sub(r"(?i)\bas noted in(?: the)? CV\b", "", text)
    text = re.sub(r"(?i)\baccording to the CV\b", "", text)
    text = re.sub(r"(?i)\bper the CV\b", "", text)
    text = re.sub(r"(?i)\bfrom the CV\b", "", text)
    text = re.sub(r"(?i)\bin the CV\b", "", text)
    text = re.sub(r"(?i)\bthe CV\b", "", text)
    text = re.sub(r"(?i)\bthis CV\b", "", text)

    text = tidy(text)
    # Fix doubled spaces / empty sentences after removals
    text = re.sub(r"\.\s*\.", ".", text)
    text = re.sub(r"^\.\s*", "", text)
    text = re.sub(r"\s+\.", ".", text)
    text = tidy(text)
    return text if text != original else original


def clean_az(text: str) -> str:
    original = text

    text = re.sub(
        r"(?i)\s*CV-də nəşrlər siyahısına istinad var, lakin istinadların özü verilməyib\.?",
        "",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də altı patent və ya müəllif şəhadətnaməsi, o cümlədən",
        "Altı patent və ya müəllif şəhadətnaməsi var, o cümlədən",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də göstərilən buraxılış işi",
        "Buraxılış işi",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də Strasburq və Bakıdakı əvvəlki təhsil və müəllim köməkçisi işi də göstərilib\.?",
        " Strasburq və Bakıda əvvəlki təhsil və müəllim köməkçisi işi də var.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də göstərilən məqalələr",
        "Məqalələr",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də h-indeks 4 kimi verilib\.?",
        " h-indeks: 4.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də PhD müdafiəsi 11 aprel 1998-ci il, habilitasiya müdafiəsi 3 iyun 2009-cu il tarixinə verilir; müəssisə və mövzu göstərilmir\.?",
        " PhD müdafiəsi: 11 aprel 1998; habilitasiya müdafiəsi: 3 iyun 2009.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də dövlət tapşırıqları və",
        "Dövlət tapşırıqları və",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də dissertasiya mövzuları:",
        "Dissertasiya mövzuları:",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də B\.S\. yazılıb\.?",
        " BSc",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də 12 dərslik və monoqrafiya daxil olmaqla 150-dən çox nəşr, 10 elmlər namizədi və 2 elmlər doktoruna rəhbərlik göstərilir\.?",
        " 12 dərslik və monoqrafiya daxil olmaqla 150-dən çox nəşr dərc edib; 10 elmlər namizədi və 2 elmlər doktoruna rəhbərlik edib.",
        text,
    )
    text = re.sub(
        r"(?i)Həsən Kalyoncu Universitetində tamamlanmış dissertasiyalar CV-dədir",
        "Həsən Kalyoncu Universitetində tamamlanmış dissertasiyalara rəhbərlik edib",
        text,
    )
    text = re.sub(
        r"(?i)Ələddin Allahverdiyev professor və doktor kimi təqdim olunur\.\s*CV-də",
        "Ələddin Allahverdiyev professor və doktordur. Vəzifələri:",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də 450-dən çox əsər, ixtira, monoqrafiya və dərslik, 7 elmlər namizədinə rəhbərlik və 60-dan çox məruzə, o cümlədən 12 sifarişli plenar məruzə qeyd olunur\.?",
        " 450-dən çox əsər, ixtira, monoqrafiya və dərslik dərc edib; 7 elmlər namizədinə rəhbərlik edib; 60-dan çox məruzə, o cümlədən 12 sifarişli plenar məruzə oxuyub.",
        text,
    )
    text = re.sub(
        r"(?i)\s*2024-cü ilin iyununa aid CV-də 450-dən çox məqalə, 41 kitab fəsli, 7 patent və 13 monoqrafiya, 30\.000-dən çox istinad və h-indeks 85 göstərilir\.?",
        " 2024-cü ilin iyununa görə: 450-dən çox məqalə, 41 kitab fəsli, 7 patent və 13 monoqrafiya; 30.000-dən çox istinad; h-indeks 85.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV bu dərəcəni alman Dr\. habil\. dərəcəsi ilə eyniləşdirir\.?",
        " Alman Dr. habil. dərəcəsinə ekvivalentdir.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV bunu alman Dr\. və ABŞ PhD dərəcəsi ilə eyniləşdirir\.?",
        " Alman Dr. və ABŞ PhD dərəcəsinə ekvivalentdir.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV bunu alman Dr\. habil\. dərəcəsi ilə eyniləşdirir\.?",
        " Alman Dr. habil. dərəcəsinə ekvivalentdir.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV bunu fizika üzrə Master of Arts ekvivalenti adlandırır\.?",
        " Fizika üzrə Master of Arts ekvivalentidir.",
        text,
    )
    text = re.sub(
        r"(?i)\s*CV-də Associated Professor yazılıb\.?",
        " Associate Professor",
        text,
    )

    # Broader AZ patterns
    text = re.sub(r"(?i)\bCV-də göstərilən\b", "", text)
    text = re.sub(r"(?i)\bCV-də qeyd olunur\b", "", text)
    text = re.sub(r"(?i)\bCV-də\s+", "", text)
    text = re.sub(r"(?i)\bCV-nin\b", "", text)
    text = re.sub(r"(?i)\bCV bunu\b", "Bu", text)
    text = re.sub(r"(?i)\bCV bu\b", "Bu", text)
    text = re.sub(r"(?i)\bCV\b(?!\s*[A-ZƏIİÖÜ])", "", text)  # bare CV token, not names

    text = tidy(text)
    text = re.sub(r"\.\s*\.", ".", text)
    text = re.sub(r"^\.\s*", "", text)
    text = tidy(text)
    return text if text != original else original


def clean_metrics_source(text: str, lang: str) -> str:
    """Replace 'CV' as a metrics Source value with a neutral label."""
    t = text.strip()
    if re.fullmatch(r"(?i)CV", t):
        return "Self-reported" if lang == "en" else "Özü tərəfindən"
    m = re.fullmatch(r"(?i)CV\s+(.+)", t)
    if m:
        return m.group(1).strip()
    return text


def process_doc(path: Path, lang: str) -> int:
    doc = Document(str(path))
    changed = 0
    cleaner = clean_en if lang == "en" else clean_az

    def apply_para(p) -> None:
        nonlocal changed
        raw = p.text
        if not raw or not raw.strip():
            return
        # Never rewrite document title line alone
        if raw.strip() in {"CURRICULUM VITAE", "ELMİ CV"}:
            return
        new = cleaner(raw)
        if new != raw:
            set_run_text(p, new)
            changed += 1

    for p in doc.paragraphs:
        apply_para(p)

    for table in doc.tables:
        hdr = [c.text.strip().lower() for c in table.rows[0].cells] if table.rows else []
        source_idx = next(
            (i for i, h in enumerate(hdr) if "source" in h or "mənbə" in h),
            None,
        )
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    raw = p.text
                    if not raw or not raw.strip():
                        continue
                    if ri > 0 and source_idx is not None and ci == source_idx:
                        new = clean_metrics_source(raw, lang)
                        if new != raw:
                            set_run_text(p, new)
                            changed += 1
                            continue
                    apply_para(p)

    if changed:
        tmp = path.with_suffix(path.suffix + ".tmpmeta")
        doc.save(str(tmp))
        try:
            tmp.replace(path)
        except OSError:
            # Destination locked (e.g. open in Word): leave .tmpmeta for retry
            print(f"LOCKED {path.name} -> kept {tmp.name}")
            return changed
    return changed


def remaining_meta(path: Path, lang: str) -> list[str]:
    doc = Document(str(path))
    hits = []
    if lang == "en":
        rx = re.compile(
            r"(?i)\b(?:the CV|this CV|in the CV|CV states|CV says|CV lists|CV gives|"
            r"CV prints|CV equates|CV calls|CV heading|named in the CV|listed in the CV|"
            r"employment record states|as noted in|the candidate\b(?!\s+of\b))"
        )
    else:
        rx = re.compile(r"(?i)\bCV-d[aəeü]|\bCV-nin\b|\bCV bunu\b|(?<![A-Za-zƏəİıÖöÜü])CV(?![A-Za-zƏəİıÖöÜü])(?!\s*[A-Z])")
    for p in doc.paragraphs:
        t = p.text.strip()
        if t and rx.search(t) and t not in {"CURRICULUM VITAE", "ELMİ CV"}:
            hits.append(t[:160])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = p.text.strip()
                    if t and rx.search(t):
                        hits.append(t[:160])
    return hits


def main() -> None:
    report = []
    total_files = 0
    total_changes = 0
    for lang in ("en", "az"):
        folder = WORD[lang]
        suffix = "_CV_EN.docx" if lang == "en" else "_CV_AZ.docx"
        for path in sorted(folder.glob(f"*{suffix}")):
            if path.name.startswith("~$"):
                continue
            n = process_doc(path, lang)
            if n:
                total_files += 1
                total_changes += n
                report.append(f"{lang} {path.name}: {n}")
    print(f"updated_files={total_files} changed_paragraphs={total_changes}")
    for line in report:
        print(line)

    # residual check
    left = []
    for lang in ("en", "az"):
        folder = WORD[lang]
        suffix = "_CV_EN.docx" if lang == "en" else "_CV_AZ.docx"
        for path in sorted(folder.glob(f"*{suffix}")):
            if path.name.startswith("~$"):
                continue
            hits = remaining_meta(path, lang)
            for h in hits:
                left.append(f"{path.name}: {h}")
    out = ROOT / "helpers" / "_cv_strip_meta_report.txt"
    out.write_text(
        f"updated_files={total_files}\nchanged_paragraphs={total_changes}\n\n"
        + "\n".join(report)
        + "\n\nREMAINING\n"
        + "\n".join(left[:80]),
        encoding="utf-8",
    )
    print(f"remaining_hits={len(left)} report={out}")


if __name__ == "__main__":
    main()
