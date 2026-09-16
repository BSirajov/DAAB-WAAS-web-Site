#!/usr/bin/env python3
"""Build the English Word file and AZ/EN HTML pages from the AZ working document."""
from __future__ import annotations

import html
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = Path(r"C:\dev\Complex Topics Clear Explanations")
AZ_DOCX = SRC_DIR / "Complex_Topics_Clear_Explanations_AZ.docx"
EN_DOCX = SRC_DIR / "Complex_Topics_Clear_Explanations_EN.docx"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _ctce_az_en_map import HEADING_IDS, MAP, TOC_TOP  # noqa: E402


def set_paragraph_text(para: Paragraph, text: str) -> None:
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_mapped(para: Paragraph) -> bool:
    key = para.text
    if key in MAP:
        set_paragraph_text(para, MAP[key])
        return True
    return False


def walk_sdt_paragraphs(doc: Document):
    for child in doc.element.body:
        if child.tag != qn("w:sdt"):
            continue
        for p_el in child.findall(".//" + qn("w:p")):
            yield Paragraph(p_el, doc)


def build_en_docx() -> list[str]:
    shutil.copyfile(AZ_DOCX, EN_DOCX)
    doc = Document(str(EN_DOCX))
    missing: list[str] = []

    def consider(para: Paragraph) -> None:
        raw = para.text
        text = raw.strip()
        if not text:
            return
        if text.isdigit() or set(text) <= {"_", " "}:
            return
        if text in MAP:
            replace_mapped(para)
            return
        if "\t" in raw:
            title, rest = raw.split("\t", 1)
            if title in MAP:
                set_paragraph_text(para, MAP[title] + "\t" + rest)
                return
        missing.append(text)

    for para in doc.paragraphs:
        consider(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    consider(para)
    for para in walk_sdt_paragraphs(doc):
        consider(para)
    doc.save(str(EN_DOCX))
    return missing


def style_name(para: Paragraph) -> str:
    try:
        return para.style.name if para.style else "Normal"
    except Exception:
        return "Normal"


def iter_body_blocks(doc: Document):
    tables = list(doc.tables)
    table_i = 0
    for child in doc.element.body:
        tag = child.tag.split("}")[-1]
        if tag == "sdt":
            continue
        if tag == "p":
            yield "p", Paragraph(child, doc)
        elif tag == "tbl":
            yield "t", tables[table_i]
            table_i += 1


def heading_id(az_text: str) -> str:
    return HEADING_IDS.get(az_text, "")


def esc(text: str) -> str:
    return html.escape(text, quote=True)


def render_table(table, lang: str) -> str:
    rows = []
    for ri, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            raw = cell.text.replace("\n", " ").strip()
            text = MAP.get(raw, raw) if lang == "en" else raw
            tag = "th" if ri == 0 else "td"
            cells.append(f"<{tag}>{esc(text)}</{tag}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<div class="ct-table-wrap"><table class="ct-doc-table">' + "".join(rows) + "</table></div>"


def render_blocks(doc: Document, lang: str) -> str:
    out: list[str] = []
    list_buf: list[str] = []
    list_kind = ""

    def flush_list() -> None:
        nonlocal list_buf, list_kind
        if not list_buf:
            return
        tag = "ol" if list_kind == "ol" else "ul"
        out.append(f"<{tag}>" + "".join(f"<li>{esc(item)}</li>" for item in list_buf) + f"</{tag}>")
        list_buf = []
        list_kind = ""

    def lang_text(az: str) -> str:
        if lang == "en":
            return MAP.get(az, az)
        return az

    skip_titles = {
        "Dünya Azərbaycanlı Alimlər Birliyi",
        "İnformatika və İKT müsabiqəsinin təşkili və iştirak qaydaları",
        "Complex Topics, Clear Explanations",
        "Mündəricat",
        "15 sentyabr 2026  •  Maliyyə məlumatları yenilənmiş versiya",
    }

    for kind, item in iter_body_blocks(doc):
        if kind == "t":
            flush_list()
            out.append(render_table(item, lang))
            continue
        para = item
        az = para.text.strip()
        if not az:
            flush_list()
            continue
        if az == "15 sentyabr 2026  •  Maliyyə məlumatları yenilənmiş versiya":
            continue
        style = style_name(para)
        if style == "Title" and az in skip_titles:
            flush_list()
            continue
        if style == "Subtitle" and az in skip_titles:
            flush_list()
            continue
        text = lang_text(az)
        if style == "List Bullet":
            if list_kind not in ("", "ul"):
                flush_list()
            list_kind = "ul"
            list_buf.append(text)
            continue
        if style == "List Number":
            if list_kind not in ("", "ol"):
                flush_list()
            list_kind = "ol"
            list_buf.append(text)
            continue
        flush_list()
        hid = heading_id(az)
        id_attr = f' id="{hid}"' if hid else ""
        if style == "Heading 1":
            out.append(f"<article class=\"ct-guide-block\"{id_attr}><h2>{esc(text)}</h2>")
            continue
        if style == "Heading 2":
            if hid in {"document-purpose", "document-status", "reader-guide"}:
                out.append(f"<article class=\"ct-guide-block\"{id_attr}><h2>{esc(text)}</h2>")
            else:
                out.append(f"<h3{id_attr}>{esc(text)}</h3>")
            continue
        if style == "Title":
            out.append(f"<h2 class=\"ct-doc-part\">{esc(text)}</h2>")
            continue
        if style == "Subtitle":
            out.append(f"<p class=\"ct-doc-kicker\">{esc(text)}</p>")
            continue
        cls = "ct-blank" if "___" in az else ""
        class_attr = f' class="{cls}"' if cls else ""
        out.append(f"<p{class_attr}>{esc(text)}</p>")

    flush_list()
    # Close leftover article tags by wrapping groups later — instead inject closing
    # tags before each new article. Post-process:
    html_out = []
    open_article = False
    for chunk in out:
        if chunk.startswith("<article"):
            if open_article:
                html_out.append("</article>")
            open_article = True
        html_out.append(chunk)
    if open_article:
        html_out.append("</article>")
    return "\n".join(html_out)


def toc_html(lang: str) -> str:
    items = []
    for hid, az_label, en_label in TOC_TOP:
        label = en_label if lang == "en" else az_label
        items.append(f'<a href="#{hid}">{esc(label)}</a>')
    aria = "On this page" if lang == "en" else "Bu səhifədə"
    return f'<nav class="ct-guide-toc" aria-label="{aria}">\n' + "\n".join(items) + "\n</nav>"


def page_note(lang: str) -> str:
    if lang == "en":
        return (
            "<p class=\"ct-guide-note\">This page presents the working approach to organising and "
            "running the competition, judging entries, and related matters. The text is not an "
            "approved regulation. Dates, scoring thresholds and procedures are initial proposals. "
            "The budget and sources of funding have not yet been decided; it is not known whether "
            "financial support or material prizes will be possible.</p>"
        )
    return (
        "<p class=\"ct-guide-note\">Bu səhifə müsabiqənin təşkili, keçirilməsi, işlərin "
        "qiymətləndirilməsi və digər əlaqəli məsələlər üzrə işçi yanaşmanı təqdim edir. Mətn "
        "təsdiqlənmiş Əsasnamə deyil. Tarixlər, qiymətləndirmə hədləri və prosedurlar ilkin "
        "təkliflərdir. Büdcə və maliyyə mənbələri hələ müəyyən edilməyib; maliyyə dəstəyinin və "
        "maddi mükafatların mümkün olub-olmayacağı məlum deyil.</p>"
    )


def splice_main(template: str, inner: str) -> str:
    start = template.find("<main")
    end = template.find("</main>")
    if start < 0 or end < 0:
        raise RuntimeError("main landmark missing")
    open_end = template.find(">", start) + 1
    return template[:open_end] + "\n" + inner + "\n" + template[end:]


def update_panel(html_text: str, lang: str) -> str:
    if lang == "en":
        old = (
            "Open competition for clear, scientifically accurate explanations of hard ICT topics "
            "— article, video, presentation or teaching pack."
        )
        new = (
            "Purpose, participation, judging and publication. "
            "Budget, funding and cash prizes have not been decided."
        )
    else:
        old = (
            "Çətin İKT mövzularının sadə və elmi cəhətdən düzgün izahı üzrə açıq müsabiqə — "
            "məqalə, video, təqdimat və ya tədris paketi."
        )
        new = (
            "Məqsəd, iştirak, qiymətləndirmə və yayım. "
            "Büdcə, maliyyələşdirmə və pul mükafatları müəyyən edilməyib."
        )
    return html_text.replace(old, new, 1)


def bump_css(html_text: str) -> str:
    return html_text.replace("daab-complex-topics.css?v=11", "daab-complex-topics.css?v=12")


def build_html() -> None:
    az_doc = Document(str(AZ_DOCX))
    for lang, path in (
        ("az", ROOT / "az" / "complex-topics.html"),
        ("en", ROOT / "en" / "complex-topics.html"),
    ):
        template = path.read_text(encoding="utf-8")
        inner = (
            page_note(lang)
            + "\n"
            + toc_html(lang)
            + '\n<section class="ct-guide" aria-label="'
            + ("Competition working document" if lang == "en" else "Müsabiqənin işçi sənədi")
            + '">\n'
            + render_blocks(az_doc, lang)
            + "\n</section>\n"
        )
        html_text = splice_main(template, inner)
        html_text = update_panel(html_text, lang)
        html_text = bump_css(html_text)
        path.write_text(html_text, encoding="utf-8")


def main() -> None:
    if not AZ_DOCX.exists():
        raise SystemExit(f"missing {AZ_DOCX}")
    missing = build_en_docx()
    leftover = [m for m in missing if m not in MAP]
    print("EN docx:", EN_DOCX)
    print("untranslated:", len(leftover))
    for item in leftover:
        print(" -", item[:160])
    build_html()
    print("HTML updated")


if __name__ == "__main__":
    main()
