#!/usr/bin/env python3
"""Extract flyer format names and tooltips into a Word document."""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from _paths import ROOT

NAVY = RGBColor(0x06, 0x31, 0x4E)
INK = RGBColor(0x1A, 0x2E, 0x3D)
MUTED = RGBColor(0x34, 0x5D, 0x76)
HEADER_FILL = "06314E"
ALT_FILL = "F3F9FD"
OUT = ROOT / "documents" / "Complex_Topics_Submission_Formats.docx"

FORMAT_RE = re.compile(
    r'<div class="ct-format"[^>]*>.*?</svg></span>'
    r'(?P<name>[^<]+)'
    r'<span class="ct-format-tip"[^>]*role="tooltip">(?P<tip>[^<]+)</span>',
    re.S,
)


def shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run(run, text: str, *, bold=False, color=INK, size=11, italic=False) -> None:
    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = color


def extract_panel(path: Path) -> dict:
    raw = path.read_text(encoding="utf-8")
    heading = re.search(r'<div class="ct-card">\s*<h3>([^<]+)</h3>\s*<p>([^<]+)</p>', raw)
    note = re.search(r'<p class="ct-format-note">([^<]+)</p>', raw)
    submit = re.search(
        r'What is submitted\?</small><strong>([^<]+)</strong>'
        r'<span class="ct-core-tip"[^>]*>([^<]+)</span>'
        r'|Nə təqdim edilir\?</small><strong>([^<]+)</strong>'
        r'<span class="ct-core-tip"[^>]*>([^<]+)</span>',
        raw,
    )
    criterion = re.search(
        r'(?:Main criterion|Əsas meyar)</small><strong>([^<]+)</strong>'
        r'<span class="ct-core-tip"[^>]*>([^<]+)</span>',
        raw,
    )
    formats = [
        (html.unescape(m.group("name").strip()), html.unescape(m.group("tip").strip()))
        for m in FORMAT_RE.finditer(raw)
    ]
    submit_label = ""
    submit_tip = ""
    if submit:
        submit_label = html.unescape((submit.group(1) or submit.group(3) or "").strip())
        submit_tip = html.unescape((submit.group(2) or submit.group(4) or "").strip())
    return {
        "heading": html.unescape(heading.group(1).strip()) if heading else "",
        "lead": html.unescape(heading.group(2).strip()) if heading else "",
        "note": html.unescape(note.group(1).strip()) if note else "",
        "submit_label": submit_label,
        "submit_tip": submit_tip,
        "criterion_label": html.unescape(criterion.group(1).strip()) if criterion else "",
        "criterion_tip": html.unescape(criterion.group(2).strip()) if criterion else "",
        "formats": formats,
    }


def add_section(doc: Document, title: str, data: dict, headers: list[str]) -> None:
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    set_run(p.add_run(data["lead"]), data["lead"], size=12)

    if data["submit_label"]:
        p = doc.add_paragraph()
        set_run(p.add_run(data["submit_label"]), data["submit_label"], bold=True, size=11)
        p = doc.add_paragraph()
        set_run(p.add_run(data["submit_tip"]), data["submit_tip"], size=11)

    if data["criterion_label"]:
        p = doc.add_paragraph()
        set_run(p.add_run(data["criterion_label"]), data["criterion_label"], bold=True, size=11)
        p = doc.add_paragraph()
        set_run(p.add_run(data["criterion_tip"]), data["criterion_tip"], size=11)

    p = doc.add_paragraph()
    set_run(p.add_run(data["note"]), data["note"], italic=True, color=MUTED, size=11)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, label in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        set_run(para.add_run(label), label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=11)
        shade_cell(hdr[i], HEADER_FILL)

    for i, (name, tip) in enumerate(data["formats"]):
        row = table.add_row().cells
        row[0].text = ""
        set_run(row[0].paragraphs[0].add_run(name), name, bold=True, size=11)
        row[1].text = ""
        set_run(row[1].paragraphs[0].add_run(tip), tip, size=11)
        if i % 2 == 1:
            shade_cell(row[0], ALT_FILL)
            shade_cell(row[1], ALT_FILL)

    doc.add_paragraph()


def main() -> None:
    az = extract_panel(ROOT / "az" / "complex-topics.html")
    en = extract_panel(ROOT / "en" / "complex-topics.html")
    if len(az["formats"]) != 9 or len(en["formats"]) != 9:
        raise SystemExit(f"expected 9 formats, got AZ={len(az['formats'])} EN={len(en['formats'])}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = INK

    title = doc.add_heading("Çətin mövzu, aydın izah — təqdimat formatları", 0)
    for run in title.runs:
        run.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    set_run(
        sub.add_run("Complex Topics, Clear Explanations — submission formats"),
        "Complex Topics, Clear Explanations — submission formats",
        italic=True,
        color=MUTED,
        size=14,
    )

    intro = doc.add_paragraph()
    set_run(
        intro.add_run(
            "This document lists the formats shown on the competition flyer panel "
            "“Öz izahını təqdim et” / “Submit your explanation”. "
            "Each explanation is the tooltip text from that panel. "
            "Recommended lengths in the tooltips are working suggestions, not confirmed limits."
        ),
        "This document lists the formats shown on the competition flyer panel "
        "“Öz izahını təqdim et” / “Submit your explanation”. "
        "Each explanation is the tooltip text from that panel. "
        "Recommended lengths in the tooltips are working suggestions, not confirmed limits.",
        size=11,
    )

    add_section(doc, "Azərbaycanca", az, ["Format", "İzah (paneldəki izah qutusu)"])
    add_section(doc, "English", en, ["Format", "Explanation (panel tooltip)"])

    footer = doc.add_paragraph()
    set_run(
        footer.add_run("Source: daab-waas.com — az/complex-topics.html and en/complex-topics.html."),
        "Source: daab-waas.com — az/complex-topics.html and en/complex-topics.html.",
        italic=True,
        color=MUTED,
        size=10,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
