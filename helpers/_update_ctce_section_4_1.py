#!/usr/bin/env python3
"""Replace section 4.1 in the AZ/EN competition Word guides with flyer-panel formats."""
from __future__ import annotations

from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from _build_complex_topics_formats_docx import extract_panel
from _paths import ROOT

AZ_HEADING = "4.1 Formatlar və həcm"
EN_HEADING = "4.1 Formats and volume"

AZ_INTRO = (
    "Flayerdəki «Öz izahını təqdim et» panelinə görə bir format seçmək, "
    "yaxud bir neçəsini birləşdirmək olar. Mövzu nə qədər mürəkkəb olsa da, "
    "məqsəd onu elmi dəqiqliyi qorumaqla aydın, məntiqli və auditoriyaya yaxın "
    "dildə izah etməkdir. Yanaşma mövzunu hədəf auditoriyaya ən yaxşı izah "
    "edən üsul olmalıdır. Aşağıdakı izahlar həmin paneldəki izah qutularındandır. "
    "Həcm göstəriciləri tövsiyədir, təsdiqlənmiş minimum və maksimum hədlər deyil."
)

EN_INTRO = (
    "The flyer panel “Submit your explanation” allows one format or a combination. "
    "However complex the topic, the aim is to explain it clearly, logically and in "
    "language close to the audience, while preserving scientific accuracy. Choose "
    "the approach that best explains the topic to the intended audience. The "
    "explanations below are the tooltips from that panel. The volumes are "
    "recommendations, not approved minimums or maximums."
)

AZ_HEADERS = ("Format", "Təklif olunan həcm", "İzah")
EN_HEADERS = ("Format", "Proposed volume", "Explanation")

AZ_VOLUME = {
    "Məqalə": "təxminən 2 000–5 000 söz",
    "Videomühazirə": "təxminən 8–20 dəqiqə",
    "Təqdimat": "təxminən 15–30 slayd",
    "Tədris paketi": "Ayrıca razılaşdırılır",
}
EN_VOLUME = {
    "Article": "about 2,000–5,000 words",
    "Video lecture": "about 8–20 minutes",
    "Presentation": "about 15–30 slides",
    "Teaching pack": "To be agreed separately",
}
AZ_VOLUME_FALLBACK = "—"
EN_VOLUME_FALLBACK = "—"


def set_run_text(paragraph, text: str, *, bold=None) -> None:
    if not paragraph.runs:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        if bold is not None:
            run.bold = bold
        return
    paragraph.runs[0].text = text
    if bold is not None:
        paragraph.runs[0].bold = bold
    for extra in paragraph.runs[1:]:
        extra.text = ""


def set_cell(cell, text: str, *, bold=False) -> None:
    para = cell.paragraphs[0]
    set_run_text(para, text, bold=bold)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def find_section_table(doc: Document, heading: str) -> tuple[Paragraph, Table]:
    intro = None
    seen = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            if para.text.strip() == heading:
                seen = True
                continue
            if seen and intro is None and para.text.strip():
                intro = para
        elif seen and child.tag == qn("w:tbl"):
            if intro is None:
                raise SystemExit(f"no intro paragraph after {heading}")
            return intro, Table(child, doc)
    raise SystemExit(f"section or table not found: {heading}")


def fill_table(table: Table, headers: tuple[str, str, str], rows: list[tuple[str, str, str]]) -> None:
    needed = 1 + len(rows)
    while len(table.rows) < needed:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > needed:
        table._tbl.remove(table.rows[-1]._tr)

    for i, text in enumerate(headers):
        set_cell(table.rows[0].cells[i], text, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            set_cell(table.rows[i].cells[j], text, bold=False)


def update_doc(path, heading: str, intro: str, headers, panel: dict, volumes: dict, fallback: str) -> None:
    doc = Document(str(path))
    intro_para, table = find_section_table(doc, heading)
    set_run_text(intro_para, intro)
    rows = [
        (name, volumes.get(name, fallback), tip)
        for name, tip in panel["formats"]
    ]
    fill_table(table, headers, rows)
    doc.save(path)
    print(f"updated {path.name}: {len(rows)} formats")


def main() -> None:
    az_panel = extract_panel(ROOT / "az" / "complex-topics.html")
    en_panel = extract_panel(ROOT / "en" / "complex-topics.html")
    if len(az_panel["formats"]) != 9 or len(en_panel["formats"]) != 9:
        raise SystemExit("expected 9 formats from the flyer panel")
    update_doc(
        ROOT / "documents" / "Complex_Topics_Clear_Explanations_AZ.docx",
        AZ_HEADING,
        AZ_INTRO,
        AZ_HEADERS,
        az_panel,
        AZ_VOLUME,
        AZ_VOLUME_FALLBACK,
    )
    update_doc(
        ROOT / "documents" / "Complex_Topics_Clear_Explanations_EN.docx",
        EN_HEADING,
        EN_INTRO,
        EN_HEADERS,
        en_panel,
        EN_VOLUME,
        EN_VOLUME_FALLBACK,
    )


if __name__ == "__main__":
    main()
