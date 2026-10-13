# -*- coding: utf-8 -*-
"""Normalize ORCID / LinkedIn / Google Scholar / website links in CV contact tables.

- Classifies values in the contact table by URL type
- Renames catch-all "ResearchGate / LinkedIn / website|vebsayt" rows to a specific label
- Expands bare ORCID IDs to https://orcid.org/...
- Adds Word hyperlinks on URL cells

Usage:
  python helpers/_normalize_cv_contact_links.py --lang az
  python helpers/_normalize_cv_contact_links.py --lang en --dry-run
  python helpers/_normalize_cv_contact_links.py --lang both
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
FOLDERS = {
    "en": ROOT / "en" / "scientists" / "CV" / "Word",
    "az": ROOT / "az" / "scientists" / "CV" / "Word",
}

ORCID_ID_RE = re.compile(r"^\d{4}-\d{4}-\d{4}-\d{3}[\dX]$", re.I)
PROFILE_LABEL_RE = re.compile(
    r"orcid|linkedin|google\s*scholar|researchgate|website|vebsayt|academia",
    re.I,
)
CATCHALL_RE = re.compile(
    r"researchgate\s*/\s*linkedin\s*/\s*(website|vebsayt)",
    re.I,
)
CONTACT_HEADING_RE = re.compile(
    r"^(contact\b|əlaqə\b|elaqe\b)|identifikasiya|identifiers",
    re.I,
)

# Internal kind -> display label per language.
LABELS = {
    "en": {
        "ORCID": "ORCID",
        "LinkedIn": "LinkedIn",
        "Google Scholar": "Google Scholar",
        "ResearchGate": "ResearchGate",
        "Website": "Website",
    },
    "az": {
        "ORCID": "ORCID",
        "LinkedIn": "LinkedIn",
        "Google Scholar": "Google Scholar",
        "ResearchGate": "ResearchGate",
        "Website": "Vebsayt",
    },
}


def is_contact_heading(text: str) -> bool:
    return bool(CONTACT_HEADING_RE.search((text or "").strip()))


def contact_table(doc: Document) -> Table | None:
    blocks: list[tuple[str, object]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("p", Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("t", Table(child, doc)))
    for i, (kind, obj) in enumerate(blocks):
        if kind == "p" and is_contact_heading(obj.text or ""):
            if i + 1 < len(blocks) and blocks[i + 1][0] == "t":
                return blocks[i + 1][1]  # type: ignore[return-value]
    return None


def normalize_url(value: str) -> str:
    v = value.strip().rstrip(".,);]")
    if ORCID_ID_RE.match(v):
        return f"https://orcid.org/{v}"
    if v.lower().startswith("orcid.org/"):
        return "https://" + v
    if v.lower().startswith("www."):
        return "https://" + v
    return v


def classify(url_or_value: str) -> str | None:
    v = normalize_url(url_or_value)
    low = v.lower()
    if "orcid.org/" in low or ORCID_ID_RE.match(url_or_value.strip()):
        return "ORCID"
    if "linkedin.com" in low:
        return "LinkedIn"
    if "scholar.google" in low:
        return "Google Scholar"
    if "researchgate.net" in low:
        return "ResearchGate"
    if low.startswith("http://") or low.startswith("https://"):
        return "Website"
    return None


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_cell_text(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    p0 = paragraphs[0]
    clear_paragraph(p0)
    run = p0.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    for extra in paragraphs[1:]:
        clear_paragraph(extra)


def set_cell_hyperlink(cell, url: str, display: str | None = None) -> None:
    display = display or url
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph("")
        paragraphs = cell.paragraphs
    p0 = paragraphs[0]
    clear_paragraph(p0)
    for extra in paragraphs[1:]:
        clear_paragraph(extra)

    part = p0.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    text = OxmlElement("w:t")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = display
    new_run.append(text)
    hyperlink.append(new_run)
    p0._p.append(hyperlink)


def process_doc(path: Path, lang: str, *, dry_run: bool) -> list[str]:
    doc = Document(str(path))
    table = contact_table(doc)
    if table is None:
        return [f"{path.name}: no contact table"]

    label_map = LABELS[lang]
    changes: list[str] = []
    for row in table.rows:
        cells = row.cells
        if len(cells) < 2:
            continue
        label = cells[0].text.strip()
        value = cells[1].text.strip()
        if not value:
            continue
        if not (
            PROFILE_LABEL_RE.search(label)
            or ORCID_ID_RE.match(value)
            or value.lower().startswith("http")
        ):
            continue

        kind = classify(value)
        if kind is None:
            continue

        new_url = normalize_url(value)
        new_label = label_map[kind]
        old_label = label
        should_retitle = bool(CATCHALL_RE.search(label)) or (
            PROFILE_LABEL_RE.search(label)
            and label.casefold() != new_label.casefold()
            and kind in label_map
        )
        if not PROFILE_LABEL_RE.search(label) and not ORCID_ID_RE.match(value):
            continue

        if should_retitle and old_label != new_label:
            changes.append(f"  label: {old_label!r} -> {new_label!r}")
            if not dry_run:
                set_cell_text(cells[0], new_label)
        elif ORCID_ID_RE.match(value) or value != new_url:
            changes.append(f"  value: {value!r} -> {new_url!r} ({new_label})")
        else:
            changes.append(f"  hyperlink: {new_label} -> {new_url}")

        if not dry_run:
            set_cell_hyperlink(cells[1], new_url, display=new_url)

    if changes and not dry_run:
        tmp = path.with_suffix(path.suffix + ".tmp.docx")
        try:
            doc.save(str(tmp))
            tmp.replace(path)
        except OSError:
            alt = path.parent / f"__tmp_{path.name}"
            doc.save(str(alt))
            alt.replace(path)
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass
    if changes:
        return [f"{path.name}:"] + changes
    return []


def run_lang(lang: str, *, dry_run: bool) -> None:
    folder = FOLDERS[lang]
    if not folder.is_dir():
        raise SystemExit(f"Missing folder: {folder}")
    reports: list[str] = []
    for path in sorted(folder.glob("*.docx")):
        if path.name.startswith("~$"):
            continue
        report = process_doc(path, lang, dry_run=dry_run)
        if report:
            reports.extend(report)
    mode = "DRY-RUN" if dry_run else "UPDATED"
    print(f"{mode} {lang}: touched {sum(1 for line in reports if line.endswith(':'))} files")
    for line in reports:
        print(line)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--lang", choices=("en", "az", "both"), default="en")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    langs = ("en", "az") if args.lang == "both" else (args.lang,)
    for lang in langs:
        run_lang(lang, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
